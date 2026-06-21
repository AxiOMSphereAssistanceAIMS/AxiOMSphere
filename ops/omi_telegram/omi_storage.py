"""
omi_storage.py
──────────────
StorageManager — все сценарии хранения для Omi.
Управляет файлами, папками, БД и миграцией.
"""

import json
import os
import re
import shutil
import sqlite3
import sys
import platform
import subprocess
import hashlib
from contextlib import contextmanager
from urllib import request as _ureq
from pathlib import Path
from datetime import datetime
from typing import Optional

try:
    from ops.report_registry_link import score_report_vs_registry_row
except ImportError:  # pragma: no cover
    from report_registry_link import score_report_vs_registry_row

try:
    from ops.omi_telegram.report_registry_semantic import (
        format_sync_note_from_ai,
        read_report_excerpt,
        semantic_match_report_to_candidates,
    )
except ImportError:  # pragma: no cover

    def read_report_excerpt(path: Path, max_chars: int = 2400) -> str:
        return ""

    def semantic_match_report_to_candidates(**_kwargs):
        return None

    def format_sync_note_from_ai(*_a, **_k):
        return ""


def _parse_db_datetime_ts(raw: str) -> float | None:
    """Парсит date_added/date_modified из SQLite (ISO-подобная строка)."""
    raw = (raw or "").strip()
    if not raw:
        return None
    for cand in (raw, raw.replace(" ", "T")):
        try:
            return datetime.fromisoformat(cand).timestamp()
        except Exception:
            continue
    return None


def _latest_registry_activity_ts(date_added: str, date_modified: str) -> float | None:
    """
    Момент «последней активности» строки в aims_registry.documents для окон «недавно».
    sync_from_ocr_registry при UPDATE не трогает date_added — остаётся старым, обновляется
    date_modified; COALESCE(date_added, date_modified) тогда скрывает свежие синки.
    """
    best: float | None = None
    for raw in (date_added, date_modified):
        ts = _parse_db_datetime_ts((raw or "").strip())
        if ts is not None and (best is None or ts > best):
            best = ts
    return best


def _file_mtime_short(file_path: str) -> str:
    """Время изменения файла на диске (если путь существует)."""
    try:
        p = Path(file_path)
        if p.is_file():
            return datetime.fromtimestamp(p.stat().st_mtime).strftime("%Y-%m-%d %H:%M")
    except Exception:
        pass
    return ""


def _file_times_iso(file_path: str) -> tuple[str, str]:
    """(created_at, mtime_at) в ISO, если файл доступен."""
    try:
        p = Path(file_path)
        if p.is_file():
            st = p.stat()
            created = datetime.fromtimestamp(getattr(st, "st_ctime", st.st_mtime)).isoformat()
            mtime = datetime.fromtimestamp(st.st_mtime).isoformat()
            return created, mtime
    except Exception:
        pass
    return "", ""


def _sanitize_filename_stem(raw: str) -> str:
    s = (raw or "").strip()
    s = re.sub(r"^tg_\d+_[0-9a-f]{8}_?", "", s, flags=re.I)
    s = re.sub(r"\s+", "_", s)
    s = re.sub(r"[^0-9A-Za-zА-Яа-яЁё._-]+", "_", s)
    s = re.sub(r"_+", "_", s).strip("._-")
    return s[:120] or "document"


def _env_flag(name: str, default: bool = False) -> bool:
    raw = os.environ.get(name, "1" if default else "0").strip().lower()
    return raw in ("1", "true", "yes", "on")

_OPS = Path(__file__).resolve().parent.parent
if str(_OPS) not in sys.path:
    sys.path.insert(0, str(_OPS))
from sqlite_helpers import sqlite_connect_wal  # noqa: E402
# Стандартные AIMS процессы
DEFAULT_PROCESSES = {
    "P01": "Asset Strategy",
    "P02": "Risk Assessment",
    "P03": "Inspection Planning",
    "P04": "Inspection Execution",
    "P05": "Maintenance Strategy",
    "P06": "Maintenance Execution",
    "P07": "SCE Performance",
    "P08": "MoC Management",
    "P09": "Reporting KPI",
    "P10": "Competence Training",
}

# Маппинг AIMS Elements → процессы
ELEMENT_TO_PROCESS = {
    **{str(e): "P01" for e in [1, 2, 3]},
    **{str(e): "P02" for e in [9, 16]},
    **{str(e): "P03" for e in [13]},
    **{str(e): "P04" for e in [15, 20]},
    **{str(e): "P05" for e in [14]},
    **{str(e): "P06" for e in [12, 11]},
    **{str(e): "P07" for e in [4, 5, 17]},
    **{str(e): "P08" for e in [6, 8]},
    **{str(e): "P09" for e in [19, 21, 22, 23]},
    **{str(e): "P10" for e in [7, 10, 18]},
}


class StorageManager:
    def __init__(self, db_path: Path, workspace: Path):
        self.db_path  = db_path
        self.workspace = workspace
        self.ocr_db_path = Path(
            os.environ.get("OCR_REGISTRY_DB", str(self.workspace / "omi_registry.db"))
        )
        self._ensure_db()

    def sync_from_ocr_registry(self, *, limit: int = 2000) -> tuple[int, int, int]:
        """
        Лёгкая инкрементальная синхронизация ocr_documents -> documents.
        Возвращает (inserted, updated, skipped). Без исключений наружу.
        """
        if not self.ocr_db_path.is_file():
            return (0, 0, 0)
        try:
            ocr = sqlite_connect_wal(self.ocr_db_path)
            ocr.row_factory = sqlite3.Row
            with self._conn() as aims:
                has_table = ocr.execute(
                    "SELECT 1 FROM sqlite_master WHERE type='table' AND name='ocr_documents'"
                ).fetchone()
                if not has_table:
                    return (0, 0, 0)
                rows = ocr.execute(
                    """
                    SELECT file_path, file_name, file_mtime, file_hash, version, source, registered_at
                    FROM ocr_documents
                    ORDER BY id DESC
                    LIMIT ?
                    """,
                    (max(1, int(limit)),),
                ).fetchall()
                rows = list(reversed(rows))
                ins = upd = sk = 0
                now = datetime.now().isoformat()
                for r in rows:
                    fp = str(r["file_path"] or "").strip()
                    fn = str(r["file_name"] or "").strip()
                    if not fp or not fn:
                        continue
                    reg_at = str(r["registered_at"] or "").strip() or now
                    created_fs, mtime_fs = _file_times_iso(fp)
                    mtime = str(r["file_mtime"] or "").strip() or mtime_fs or reg_at
                    created = created_fs or reg_at
                    h = str(r["file_hash"] or "").strip()
                    ver = int(r["version"] or 1)
                    src = str(r["source"] or "").strip() or "ocr-sync"
                    title = fn.rsplit(".", 1)[0] if "." in fn else fn
                    ftype = fn.rsplit(".", 1)[-1].lower() if "." in fn else ""
                    notes_obj = {
                        "ocr_sync": True,
                        "file_hash": h,
                        "version": ver,
                        "registered_at": reg_at,
                    }
                    notes = json.dumps(notes_obj, ensure_ascii=False)
                    ex = aims.execute(
                        "SELECT id, notes FROM documents WHERE file_path = ?",
                        (fp,),
                    ).fetchone()
                    if ex:
                        try:
                            old = json.loads(ex["notes"] or "{}")
                            if isinstance(old, dict) and old.get("file_hash") == h and h:
                                sk += 1
                                continue
                        except Exception:
                            pass
                        aims.execute(
                            """
                            UPDATE documents
                               SET file_name = ?, file_type = ?, title = ?, source = ?,
                                   date_modified = ?, file_created_at = ?, file_mtime_at = ?,
                                   canonical_file_name = COALESCE(canonical_file_name, file_name),
                                   original_file_name = COALESCE(original_file_name, file_name),
                                   notes = ?
                             WHERE id = ?
                            """,
                            (fn, ftype, title, src, mtime, created, mtime, notes, ex["id"]),
                        )
                        upd += 1
                    else:
                        aims.execute(
                            """
                            INSERT INTO documents (
                                file_path, file_name, file_type, title, summary, aims_process,
                                is_master, is_anonymized, language, date_added, date_modified,
                                file_created_at, file_mtime_at, original_file_name, canonical_file_name,
                                source_filename, stored_path, process_code, anonymized_result_path,
                                source, notes
                            ) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
                            """,
                            (
                                fp,
                                fn,
                                ftype,
                                title,
                                None,
                                None,
                                0,
                                0,
                                "en",
                                reg_at,
                                mtime,
                                created,
                                mtime,
                                fn,
                                fn,
                                fn,
                                fp,
                                "",
                                "",
                                src,
                                notes,
                            ),
                        )
                        ins += 1
                aims.commit()
                if ins or upd:
                    self._log("sync_from_ocr_registry", detail=f"ins={ins},upd={upd},sk={sk}")
                return (ins, upd, sk)
        except Exception as _sync_err:
            print(f"[omi_storage] sync_from_ocr_registry failed: {_sync_err}", flush=True)
            return (0, 0, 0)
        finally:
            try:
                ocr.close()
            except Exception:
                pass

    # ── DB helpers ────────────────────────────────────────────

    @contextmanager
    def _conn(self):
        """Context manager: opens WAL connection, commits on success, always closes."""
        conn = sqlite_connect_wal(self.db_path)
        conn.row_factory = sqlite3.Row
        try:
            yield conn
            conn.commit()
        except Exception:
            conn.rollback()
            raise
        finally:
            conn.close()

    @staticmethod
    def _table_columns(conn: sqlite3.Connection, table: str) -> set[str]:
        rows = conn.execute(f"PRAGMA table_info({table})").fetchall()
        return {str(r[1]) for r in rows}

    def _migrate_documents_columns(self, conn: sqlite3.Connection) -> None:
        """Старая БД могла иметь documents без части колонок — добавляем без потери данных."""
        if not conn.execute(
            "SELECT 1 FROM sqlite_master WHERE type='table' AND name='documents'"
        ).fetchone():
            return
        existing = self._table_columns(conn, "documents")
        for col, decl in (
            ("file_name", "TEXT"),
            ("file_type", "TEXT"),
            ("aims_process", "TEXT"),
            ("aims_element", "TEXT"),
            ("iso_clause", "TEXT"),
            ("title", "TEXT"),
            ("keywords", "TEXT"),
            ("summary", "TEXT"),
            ("is_master", "INTEGER DEFAULT 0"),
            ("is_anonymized", "INTEGER DEFAULT 0"),
            ("language", "TEXT DEFAULT 'en'"),
            ("date_added", "TEXT"),
            ("date_modified", "TEXT"),
            ("file_created_at", "TEXT"),
            ("file_mtime_at", "TEXT"),
            ("original_file_name", "TEXT"),
            ("canonical_file_name", "TEXT"),
            ("source", "TEXT"),
            ("notes", "TEXT"),
            ("source_filename", "TEXT"),
            ("stored_path", "TEXT"),
            ("process_code", "TEXT"),
            ("anonymized_result_path", "TEXT"),
            ("master_doc_json", "TEXT"),
            ("content_hash", "TEXT"),
            ("created_at", "TEXT"),
            ("updated_at", "TEXT"),
        ):
            if col not in existing:
                conn.execute(f"ALTER TABLE documents ADD COLUMN {col} {decl}")

    def _ensure_db(self):
        self.db_path.parent.mkdir(parents=True, exist_ok=True)
        with self._conn() as conn:
            conn.executescript("""
                CREATE TABLE IF NOT EXISTS documents (
                    id            INTEGER PRIMARY KEY AUTOINCREMENT,
                    file_path     TEXT NOT NULL UNIQUE,
                    file_name     TEXT NOT NULL,
                    file_type     TEXT,
                    aims_process  TEXT,
                    aims_element  TEXT,
                    iso_clause    TEXT,
                    title         TEXT,
                    keywords      TEXT,
                    summary       TEXT,
                    is_master     INTEGER DEFAULT 0,
                    is_anonymized INTEGER DEFAULT 0,
                    language      TEXT DEFAULT 'en',
                    date_added    TEXT,
                    date_modified TEXT,
                    file_created_at TEXT,
                    file_mtime_at TEXT,
                    original_file_name TEXT,
                    canonical_file_name TEXT,
                    source        TEXT,
                    notes         TEXT,
                    master_doc_json TEXT,
                    content_hash  TEXT,
                    created_at    TEXT,
                    updated_at    TEXT
                );
                CREATE TABLE IF NOT EXISTS omi_processes (
                    code    TEXT PRIMARY KEY,
                    name    TEXT NOT NULL,
                    path    TEXT,
                    created TEXT
                );
                CREATE TABLE IF NOT EXISTS omi_config (
                    key     TEXT PRIMARY KEY,
                    value   TEXT,
                    updated TEXT
                );
                CREATE TABLE IF NOT EXISTS omi_log (
                    id          INTEGER PRIMARY KEY AUTOINCREMENT,
                    timestamp   TEXT NOT NULL,
                    action      TEXT NOT NULL,
                    file_id     INTEGER,
                    actor       TEXT,
                    detail      TEXT,
                    FOREIGN KEY (file_id) REFERENCES documents(id)
                );
                CREATE TABLE IF NOT EXISTS omi_tasks (
                    id          INTEGER PRIMARY KEY AUTOINCREMENT,
                    created_at  TEXT NOT NULL,
                    source      TEXT NOT NULL DEFAULT 'manual',
                    status      TEXT NOT NULL DEFAULT 'pending',
                    title       TEXT NOT NULL,
                    detail      TEXT,
                    chat_hint   TEXT
                );
                CREATE TABLE IF NOT EXISTS omi_skills (
                    id          INTEGER PRIMARY KEY AUTOINCREMENT,
                    name        TEXT NOT NULL UNIQUE,
                    body        TEXT NOT NULL,
                    enabled     INTEGER NOT NULL DEFAULT 1,
                    created_at  TEXT NOT NULL,
                    updated_at  TEXT NOT NULL
                );
            """)
            self._migrate_documents_columns(conn)
            conn.execute(
                "CREATE UNIQUE INDEX IF NOT EXISTS documents_content_hash"
                " ON documents(content_hash) WHERE content_hash IS NOT NULL"
            )
            self._migrate_omi_tasks_table(conn)
            self._migrate_omi_skills_table(conn)
            self._migrate_night_plan_columns(conn)
            self._migrate_fts_with_filename(conn)
            self._migrate_vector_store(conn)
            self._migrate_solution_memory(conn)
            conn.executescript("""
                CREATE TABLE IF NOT EXISTS omi_chat_context (
                    chat_id          INTEGER PRIMARY KEY,
                    last_file_name   TEXT,
                    last_file_path   TEXT,
                    last_process     TEXT,
                    last_doc_type    TEXT,
                    last_search      TEXT,
                    last_doc_id      INTEGER,
                    updated_at       TEXT NOT NULL
                );
                CREATE INDEX IF NOT EXISTS idx_process  ON documents(aims_process);
                CREATE INDEX IF NOT EXISTS idx_keywords ON documents(keywords);
                CREATE INDEX IF NOT EXISTS idx_omi_tasks_status ON omi_tasks(status);
                CREATE INDEX IF NOT EXISTS idx_omi_tasks_night ON omi_tasks(scheduled_for, status);
                CREATE INDEX IF NOT EXISTS idx_omi_skills_enabled ON omi_skills(enabled);
            """)
            # Seed default processes
            for code, name in DEFAULT_PROCESSES.items():
                conn.execute(
                    "INSERT OR IGNORE INTO omi_processes(code,name,path,created) VALUES(?,?,?,?)",
                    (code, name, str(self.workspace / f"master/{code}_{name.replace(' ','_')}"),
                     datetime.now().isoformat())
                )

    def _log(self, action: str, detail: str = "", file_id: int | None = None, actor: str | None = None):
        try:
            with self._conn() as conn:
                conn.execute(
                    """
                    INSERT INTO omi_log (timestamp, action, file_id, actor, detail)
                    VALUES (?,?,?,?,?)
                    """,
                    (
                        datetime.now().isoformat(),
                        action,
                        file_id,
                        actor,
                        detail[:4000] if detail else "",
                    ),
                )
                conn.commit()
        except Exception:
            # старые БД без omi_log — не ломаем основной поток
            pass

    def _migrate_omi_tasks_table(self, conn: sqlite3.Connection) -> None:
        """Очередь задач для Omi (Axi → оператор / структура БД)."""
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS omi_tasks (
                id          INTEGER PRIMARY KEY AUTOINCREMENT,
                created_at  TEXT NOT NULL,
                source      TEXT NOT NULL DEFAULT 'manual',
                status      TEXT NOT NULL DEFAULT 'pending',
                title       TEXT NOT NULL,
                detail      TEXT,
                chat_hint   TEXT
            );
            """
        )

    def _migrate_omi_skills_table(self, conn: sqlite3.Connection) -> None:
        """Черновики skills для Omi-LLM (только владелец редактирует)."""
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS omi_skills (
                id          INTEGER PRIMARY KEY AUTOINCREMENT,
                name        TEXT NOT NULL UNIQUE,
                body        TEXT NOT NULL,
                enabled     INTEGER NOT NULL DEFAULT 1,
                created_at  TEXT NOT NULL,
                updated_at  TEXT NOT NULL
            );
            """
        )

    def enqueue_task(
        self,
        title: str,
        detail: str = "",
        *,
        source: str = "manual",
        chat_hint: str | None = None,
    ) -> str:
        title = (title or "").strip()
        if not title:
            return "❌ Пустой заголовок задачи."
        now = datetime.now().isoformat()
        with self._conn() as conn:
            conn.execute(
                """
                INSERT INTO omi_tasks (created_at, source, status, title, detail, chat_hint)
                VALUES (?,?,?,?,?,?)
                """,
                (now, (source or "manual")[:32], "pending", title[:500], (detail or "")[:8000], chat_hint),
            )
            conn.commit()
            tid = conn.execute("SELECT last_insert_rowid()").fetchone()[0]
        self._log("task_enqueue", detail=f"#{tid} {title[:200]}", actor=source)
        return f"✅ Задача #{tid} добавлена в очередь."

    def list_tasks(self, status: str = "pending", limit: int = 15) -> str:
        st = (status or "pending").strip().lower()
        if st not in ("pending", "done", "all"):
            st = "pending"
        with self._conn() as conn:
            if st == "all":
                rows = conn.execute(
                    """
                    SELECT id, created_at, source, status, title
                    FROM omi_tasks
                    ORDER BY id DESC
                    LIMIT ?
                    """,
                    (limit,),
                ).fetchall()
            else:
                rows = conn.execute(
                    """
                    SELECT id, created_at, source, status, title
                    FROM omi_tasks
                    WHERE status = ?
                    ORDER BY id ASC
                    LIMIT ?
                    """,
                    (st, limit),
                ).fetchall()
        if not rows:
            return "📋 *Очередь задач пуста.*"
        lines = ["📋 *Задачи Omi:*\n"]
        for r in rows:
            lines.append(
                f"  • `#{r['id']}` [{r['status']}] _{r['source']}_ — {r['title'][:120]}"
            )
        return "\n".join(lines)

    def complete_task(self, task_id: int) -> str:
        try:
            tid = int(task_id)
        except (TypeError, ValueError):
            return "❌ Некорректный номер задачи."
        with self._conn() as conn:
            cur = conn.execute(
                "UPDATE omi_tasks SET status = 'done' WHERE id = ? AND status = 'pending'",
                (tid,),
            )
            conn.commit()
            if cur.rowcount == 0:
                return f"ℹ️ Задача #{tid} не найдена или уже закрыта."
        self._log("task_done", detail=f"#{tid}")
        return f"✅ Задача #{tid} отмечена выполненной."

    # ── Night plan ───────────────────────────────────────────────

    def _migrate_vector_store(self, conn: sqlite3.Connection) -> None:
        """SQLite-backed vector store for local cosine-similarity search (no external deps)."""
        conn.executescript("""
            CREATE TABLE IF NOT EXISTS document_embeddings (
                doc_id   INTEGER NOT NULL,
                model    TEXT    NOT NULL,
                dim      INTEGER NOT NULL,
                vector   BLOB    NOT NULL,
                PRIMARY KEY (doc_id, model)
            );
            CREATE TABLE IF NOT EXISTS embed_models (
                model  TEXT PRIMARY KEY,
                dim    INTEGER NOT NULL,
                note   TEXT
            );
        """)

    def _migrate_solution_memory(self, conn: sqlite3.Connection) -> None:
        """Episodic solution memory: task → plan → result → feedback → skill promotion."""
        conn.executescript("""
            CREATE TABLE IF NOT EXISTS solution_memory (
                id              INTEGER PRIMARY KEY AUTOINCREMENT,
                chat_id         INTEGER,
                task_text       TEXT NOT NULL,
                plan_json       TEXT NOT NULL,
                result_summary  TEXT,
                feedback_score  INTEGER,        -- NULL=pending, 1=bad, 3=ok, 5=great
                feedback_note   TEXT,
                promoted        INTEGER NOT NULL DEFAULT 0,
                skill_id        INTEGER,
                created_at      TEXT NOT NULL,
                feedback_at     TEXT
            );
            CREATE INDEX IF NOT EXISTS idx_solution_memory_chat
                ON solution_memory(chat_id, created_at DESC);
            CREATE INDEX IF NOT EXISTS idx_solution_memory_score
                ON solution_memory(feedback_score, promoted);
        """)

    def log_solution(
        self,
        task_text: str,
        plan_json: str,
        result_summary: str = "",
        chat_id: int | None = None,
    ) -> int:
        """Save a completed multi-step execution. Returns solution id."""
        now = datetime.now().isoformat()
        with self._conn() as conn:
            cur = conn.execute(
                """
                INSERT INTO solution_memory
                    (chat_id, task_text, plan_json, result_summary, created_at)
                VALUES (?, ?, ?, ?, ?)
                """,
                (chat_id, (task_text or "")[:2000], plan_json[:8000],
                 (result_summary or "")[:2000], now),
            )
            conn.commit()
            return cur.lastrowid  # type: ignore[return-value]

    def update_solution_feedback(
        self,
        solution_id: int,
        score: int,
        note: str = "",
    ) -> bool:
        """Record user feedback (score 1-5) for a logged solution.

        Returns True if a row was found and updated, False if solution_id
        doesn't exist in solution_memory (e.g. it's a task-queue ID).
        """
        now = datetime.now().isoformat()
        with self._conn() as conn:
            cur = conn.execute(
                """
                UPDATE solution_memory
                SET feedback_score = ?, feedback_note = ?, feedback_at = ?
                WHERE id = ?
                """,
                (max(1, min(5, score)), (note or "")[:1000], now, solution_id),
            )
            conn.commit()
            found = cur.rowcount > 0
        self._log("solution_feedback", detail=f"id={solution_id} score={score} note={note[:200]} found={found}")
        return found

    def promote_solution_to_skill(self, solution_id: int) -> str:
        """Convert a high-rated solution into a reusable skill (few-shot example)."""
        with self._conn() as conn:
            row = conn.execute(
                "SELECT task_text, plan_json, result_summary, feedback_note FROM solution_memory WHERE id = ?",
                (solution_id,),
            ).fetchone()
            if not row:
                return f"⚠️ Решение #{solution_id} не найдено."
            task_text, plan_json, result_summary, feedback_note = row
            # Build skill body as a few-shot example
            try:
                import json as _json
                steps = _json.loads(plan_json)
                steps_text = "\n".join(
                    f"  {i+1}. {s.get('action')}  {_json.dumps(s.get('params', {}), ensure_ascii=False)}"
                    for i, s in enumerate(steps)
                    if isinstance(s, dict)
                )
            except Exception:
                steps_text = plan_json[:400]

            skill_name = f"solution_{solution_id}"
            skill_body = (
                f"# Пример решения задачи (из обратной связи)\n\n"
                f"**Запрос:** {task_text}\n\n"
                f"**Шаги:**\n{steps_text}\n\n"
                f"**Результат:** {result_summary or '—'}\n"
                + (f"\n**Оценка пользователя:** {feedback_note}" if feedback_note else "")
                + "\n\n_Используй этот паттерн для аналогичных запросов._"
            )
            # Save as skill
            now = datetime.now().isoformat()
            row2 = conn.execute("SELECT id FROM omi_skills WHERE name = ?", (skill_name,)).fetchone()
            if row2:
                conn.execute(
                    "UPDATE omi_skills SET body = ?, updated_at = ?, enabled = 1 WHERE name = ?",
                    (skill_body, now, skill_name),
                )
                skill_id = row2[0]
            else:
                cur = conn.execute(
                    "INSERT INTO omi_skills (name, body, enabled, created_at, updated_at) VALUES (?,?,1,?,?)",
                    (skill_name, skill_body, now, now),
                )
                skill_id = cur.lastrowid
            conn.execute(
                "UPDATE solution_memory SET promoted = 1, skill_id = ? WHERE id = ?",
                (skill_id, solution_id),
            )
            conn.commit()
        self._log("solution_promoted", detail=f"id={solution_id} skill={skill_name}")
        return f"✅ Решение #{solution_id} сохранено как skill `{skill_name}` — будет подмешиваться в контекст похожих задач."

    def get_similar_solutions(self, task_text: str, limit: int = 3) -> str:
        """Find past solutions similar to task_text by keyword overlap. Used for context injection."""
        try:
            words = set(w.lower() for w in re.split(r"\W+", task_text) if len(w) > 3)
            with self._conn() as conn:
                rows = conn.execute(
                    """
                    SELECT id, task_text, plan_json, result_summary, feedback_score
                    FROM solution_memory
                    WHERE feedback_score >= 3
                    ORDER BY created_at DESC
                    LIMIT 100
                    """,
                ).fetchall()
            if not rows:
                return ""
            scored: list[tuple[int, tuple]] = []
            for row in rows:
                row_words = set(w.lower() for w in re.split(r"\W+", row[1] or "") if len(w) > 3)
                overlap = len(words & row_words)
                if overlap > 0:
                    scored.append((overlap, row))
            scored.sort(key=lambda x: x[0], reverse=True)
            top = [r for _, r in scored[:limit]]
            if not top:
                return ""
            import json as _json
            lines = ["[Похожие успешные решения из памяти:"]
            for sol_id, task, plan_j, result, score in top:
                try:
                    steps = _json.loads(plan_j)
                    steps_short = " → ".join(
                        s.get("action", "?") for s in steps if isinstance(s, dict)
                    )
                except Exception:
                    steps_short = plan_j[:80]
                lines.append(f"• Задача: {task[:120]}")
                lines.append(f"  Шаги: {steps_short}")
                if result:
                    lines.append(f"  Итог: {result[:120]}")
                lines.append("")
            lines.append("]")
            return "\n".join(lines)
        except Exception:
            return ""

    def list_solutions(self, limit: int = 20, only_pending: bool = False) -> str:
        """List solution memory entries for the user."""
        try:
            with self._conn() as conn:
                if only_pending:
                    rows = conn.execute(
                        "SELECT id, task_text, feedback_score, promoted, created_at FROM solution_memory "
                        "WHERE feedback_score IS NULL ORDER BY created_at DESC LIMIT ?",
                        (limit,),
                    ).fetchall()
                else:
                    rows = conn.execute(
                        "SELECT id, task_text, feedback_score, promoted, created_at FROM solution_memory "
                        "ORDER BY created_at DESC LIMIT ?",
                        (limit,),
                    ).fetchall()
            if not rows:
                return "📭 Память решений пуста." if not only_pending else "✅ Все решения оценены."
            lines = [f"📚 *Память решений* ({len(rows)} записей):\n"]
            score_icon = {None: "⏳", 1: "👎", 2: "👎", 3: "👍", 4: "👍", 5: "⭐"}
            for sol_id, task, score, promoted, created in rows:
                icon = score_icon.get(score, "?")
                promo = " 🔖skill" if promoted else ""
                lines.append(f"  #{sol_id} {icon}{promo}  `{(task or '')[:80]}`")
                lines.append(f"       {created[:16]}")
            return "\n".join(lines)
        except Exception as e:
            return f"⚠️ {e}"

    def _migrate_night_plan_columns(self, conn: sqlite3.Connection) -> None:
        """Add scheduled_for and estimate_min to omi_tasks (idempotent)."""
        for col, defn in [
            ("scheduled_for", "TEXT DEFAULT NULL"),
            ("estimate_min",  "INTEGER DEFAULT NULL"),
        ]:
            try:
                conn.execute(f"ALTER TABLE omi_tasks ADD COLUMN {col} {defn}")
            except Exception:
                pass

    def _migrate_fts_with_filename(self, conn: sqlite3.Connection) -> None:
        """Rebuild documents_fts to include file_name + canonical_file_name, add sync triggers.

        Idempotent: skips if FTS schema already contains 'file_name'.
        On rebuild: drops old FTS + triggers, creates new external-content FTS5,
        re-indexes all existing documents, adds INSERT/UPDATE/DELETE triggers.
        """
        import logging as _logging
        _log = _logging.getLogger("omi.storage.fts")

        row = conn.execute(
            "SELECT sql FROM sqlite_master WHERE type='table' AND name='documents_fts'"
        ).fetchone()
        if row and "file_name" in (row[0] or ""):
            # Already migrated — ensure triggers exist (idempotent)
            self._ensure_fts_triggers(conn)
            return

        _log.info("fts_migrate: rebuilding documents_fts with file_name + canonical_file_name")
        # Drop old FTS (cascade drops shadow tables automatically)
        conn.executescript("""
            DROP TABLE IF EXISTS documents_fts;
            DROP TRIGGER IF EXISTS docs_ai_fts;
            DROP TRIGGER IF EXISTS docs_au_fts;
            DROP TRIGGER IF EXISTS docs_ad_fts;
        """)
        # Create new external-content FTS5 — index points to documents table,
        # no duplicate content storage, tokenize with unicode + diacritics removal.
        conn.executescript("""
            CREATE VIRTUAL TABLE documents_fts USING fts5(
                file_name,
                canonical_file_name,
                source_filename,
                title,
                keywords,
                summary,
                content='documents',
                content_rowid='id',
                tokenize='unicode61 remove_diacritics 2'
            );
        """)
        # Populate from existing documents
        conn.execute("""
            INSERT INTO documents_fts(rowid, file_name, canonical_file_name,
                source_filename, title, keywords, summary)
            SELECT id,
                COALESCE(file_name, ''),
                COALESCE(canonical_file_name, ''),
                COALESCE(source_filename, ''),
                COALESCE(title, ''),
                COALESCE(keywords, ''),
                COALESCE(summary, '')
            FROM documents
        """)
        conn.execute("INSERT INTO documents_fts(documents_fts) VALUES('optimize')")
        self._ensure_fts_triggers(conn)
        conn.commit()
        _log.info("fts_migrate: done — FTS rebuilt with file_name, triggers installed")

    def _ensure_fts_triggers(self, conn: sqlite3.Connection) -> None:
        """Create INSERT / UPDATE / DELETE triggers to keep documents_fts in sync."""
        existing = {r[0] for r in conn.execute(
            "SELECT name FROM sqlite_master WHERE type='trigger' AND name LIKE 'docs_%_fts'"
        ).fetchall()}
        if "docs_ai_fts" not in existing:
            conn.execute("""
                CREATE TRIGGER docs_ai_fts AFTER INSERT ON documents BEGIN
                    INSERT INTO documents_fts(rowid, file_name, canonical_file_name,
                        source_filename, title, keywords, summary)
                    VALUES (new.id,
                        COALESCE(new.file_name, ''),
                        COALESCE(new.canonical_file_name, ''),
                        COALESCE(new.source_filename, ''),
                        COALESCE(new.title, ''),
                        COALESCE(new.keywords, ''),
                        COALESCE(new.summary, ''));
                END
            """)
        if "docs_au_fts" not in existing:
            conn.execute("""
                CREATE TRIGGER docs_au_fts AFTER UPDATE ON documents BEGIN
                    INSERT INTO documents_fts(documents_fts, rowid, file_name, canonical_file_name,
                        source_filename, title, keywords, summary)
                    VALUES ('delete', old.id,
                        COALESCE(old.file_name, ''),
                        COALESCE(old.canonical_file_name, ''),
                        COALESCE(old.source_filename, ''),
                        COALESCE(old.title, ''),
                        COALESCE(old.keywords, ''),
                        COALESCE(old.summary, ''));
                    INSERT INTO documents_fts(rowid, file_name, canonical_file_name,
                        source_filename, title, keywords, summary)
                    VALUES (new.id,
                        COALESCE(new.file_name, ''),
                        COALESCE(new.canonical_file_name, ''),
                        COALESCE(new.source_filename, ''),
                        COALESCE(new.title, ''),
                        COALESCE(new.keywords, ''),
                        COALESCE(new.summary, ''));
                END
            """)
        if "docs_ad_fts" not in existing:
            conn.execute("""
                CREATE TRIGGER docs_ad_fts AFTER DELETE ON documents BEGIN
                    INSERT INTO documents_fts(documents_fts, rowid, file_name, canonical_file_name,
                        source_filename, title, keywords, summary)
                    VALUES ('delete', old.id,
                        COALESCE(old.file_name, ''),
                        COALESCE(old.canonical_file_name, ''),
                        COALESCE(old.source_filename, ''),
                        COALESCE(old.title, ''),
                        COALESCE(old.keywords, ''),
                        COALESCE(old.summary, ''));
                END
            """)
        conn.commit()

    def enqueue_night_task(
        self,
        title: str,
        detail: str = "",
        *,
        estimate_min: int | None = None,
        source: str = "user",
    ) -> str:
        title = (title or "").strip()
        if not title:
            return "❌ Пустой заголовок задачи."
        now = datetime.now().isoformat()
        with self._conn() as conn:
            conn.execute(
                """
                INSERT INTO omi_tasks
                  (created_at, source, status, title, detail, scheduled_for, estimate_min)
                VALUES (?,?,?,?,?,?,?)
                """,
                (now, source[:32], "pending", title[:500], (detail or "")[:4000],
                 "night", estimate_min),
            )
            conn.commit()
            tid = conn.execute("SELECT last_insert_rowid()").fetchone()[0]
        self._log("night_task_enqueue", detail=f"#{tid} {title[:200]}", actor=source)
        return f"✅ Задача *#{tid}* добавлена в ночной план."

    def get_night_queue_stats(self) -> dict:
        """Live counts for the night work queue."""
        stats: dict = {
            "income_files": 0,
            "omi_no_summary": 0,
            "aims_no_summary": 0,
            "custom_tasks": [],
        }
        # Income files waiting for OCR
        income_dir = self.workspace / "inbox" / "income"
        if income_dir.exists():
            stats["income_files"] = sum(1 for f in income_dir.iterdir() if f.is_file())

        # Omi-registry docs without AI summary
        if self.ocr_db_path.is_file():
            try:
                oc = sqlite_connect_wal(self.ocr_db_path)
                oc.row_factory = sqlite3.Row
                row = oc.execute(
                    "SELECT COUNT(*) FROM documents WHERE omi_ai_summary IS NULL OR omi_ai_summary=''"
                ).fetchone()
                stats["omi_no_summary"] = row[0] if row else 0
                oc.close()
            except Exception:
                pass

        # AIMS-registry docs without summary
        with self._conn() as conn:
            try:
                row = conn.execute(
                    "SELECT COUNT(*) FROM documents WHERE summary IS NULL OR summary=''"
                ).fetchone()
                stats["aims_no_summary"] = row[0] if row else 0
            except Exception:
                pass
            # Custom user night tasks
            try:
                rows = conn.execute(
                    """
                    SELECT id, title, estimate_min FROM omi_tasks
                    WHERE scheduled_for='night' AND status='pending'
                    ORDER BY id
                    """
                ).fetchall()
                stats["custom_tasks"] = [(r["id"], r["title"], r["estimate_min"]) for r in rows]
            except Exception:
                stats["custom_tasks"] = []

        return stats

    def format_night_plan(self) -> str:
        """Formatted night work plan report (Markdown)."""
        s = self.get_night_queue_stats()
        income   = s["income_files"]
        omi_sum  = s["omi_no_summary"]
        aims_sum = s["aims_no_summary"]
        custom   = s["custom_tasks"]

        # Time estimates (minutes)
        income_min   = income * 5
        omi_lo, omi_hi     = omi_sum * 1, omi_sum * 2
        aims_lo, aims_hi   = aims_sum * 1, aims_sum * 2
        maint_min    = 30
        custom_min   = sum((t[2] or 15) for t in custom)

        tot_lo = income_min + omi_lo  + aims_lo  + maint_min + custom_min
        tot_hi = income_min + omi_hi + aims_hi + maint_min + custom_min

        def _fmt(mins: int) -> str:
            h, m = divmod(int(mins), 60)
            return f"{h}ч {m:02d}м" if h else f"{m}м"

        lines = ["🌙 *Ночной план работы* — старт 01:00 Dubai\n", "```"]
        lines.append(f"{'Задача':<32} {'Кол-во':<8} {'Оценка'}")
        lines.append("─" * 55)

        def row(task, count, est_str):
            cnt = str(count) if count else "—"
            return f"{task:<32} {cnt:<8} {est_str}"

        if income > 0:
            lines.append(row("OCR (income/ → база)", income, f"~{_fmt(income_min)}"))
        else:
            lines.append(row("OCR (income/ → база)", 0, "нет файлов"))

        lines.append(row("AI summary (omi-register)", omi_sum,
                         f"~{_fmt(omi_lo)}–{_fmt(omi_hi)}" if omi_sum else "—"))
        lines.append(row("AIMS summary", aims_sum,
                         f"~{_fmt(aims_lo)}–{_fmt(aims_hi)}" if aims_sum else "—"))
        lines.append(row("Backup + maintenance", "—", f"~{_fmt(maint_min)}"))

        if custom:
            lines.append("─" * 55)
            for tid, title, est in custom:
                est_str = f"~{est}м" if est else "?"
                lines.append(f"  #{tid} {title[:38]:<38} {est_str}")

        lines.append("─" * 55)
        lines.append(row("ИТОГО", "", f"~{_fmt(tot_lo)}–{_fmt(tot_hi)}"))
        lines.append("```")

        if income == 0 and omi_sum == 0 and aims_sum == 0 and not custom:
            lines.append("\n✅ Очередь пуста. Ночью нечего делать.")
        else:
            lines.append(f"\n⏰ *01:00* — ночной режим запустится автоматически.")
            if custom:
                lines.append(f"📋 Пользовательских задач в плане: *{len(custom)}*")

        return "\n".join(lines)

    # ── Skills (Omi-LLM), только владелец — CRUD ────────────────

    def skill_upsert(self, name: str, body: str) -> str:
        name = (name or "").strip()
        body = (body or "").strip()
        if not name or not body:
            return "❌ Укажи `name` и непустой `body` skill."
        if len(name) > 200:
            return "❌ Имя skill слишком длинное."
        now = datetime.now().isoformat()
        with self._conn() as conn:
            row = conn.execute("SELECT id FROM omi_skills WHERE name = ?", (name,)).fetchone()
            if row:
                conn.execute(
                    """
                    UPDATE omi_skills SET body = ?, updated_at = ?, enabled = 1
                    WHERE name = ?
                    """,
                    (body[:120000], now, name),
                )
                msg = f"✅ Skill `{name}` обновлён."
            else:
                conn.execute(
                    """
                    INSERT INTO omi_skills (name, body, enabled, created_at, updated_at)
                    VALUES (?,?,1,?,?)
                    """,
                    (name, body[:120000], now, now),
                )
                msg = f"✅ Skill `{name}` добавлен."
            conn.commit()
        self._log("skill_upsert", detail=name[:200])
        return msg

    def skill_delete(self, name_or_id: str) -> str:
        raw = (name_or_id or "").strip()
        if not raw:
            return "❌ Укажи имя или id skill."
        with self._conn() as conn:
            if raw.isdigit():
                cur = conn.execute("DELETE FROM omi_skills WHERE id = ?", (int(raw),))
            else:
                cur = conn.execute("DELETE FROM omi_skills WHERE name = ?", (raw,))
            conn.commit()
            n = cur.rowcount
        if not n:
            return f"ℹ️ Skill `{raw}` не найден."
        self._log("skill_delete", detail=raw[:200])
        return f"✅ Удалено ({n})."

    def skill_set_enabled(self, name_or_id: str, enabled: bool) -> str:
        raw = (name_or_id or "").strip()
        en = 1 if enabled else 0
        with self._conn() as conn:
            if raw.isdigit():
                cur = conn.execute(
                    "UPDATE omi_skills SET enabled = ?, updated_at = ? WHERE id = ?",
                    (en, datetime.now().isoformat(), int(raw)),
                )
            else:
                cur = conn.execute(
                    "UPDATE omi_skills SET enabled = ?, updated_at = ? WHERE name = ?",
                    (en, datetime.now().isoformat(), raw),
                )
            conn.commit()
            n = cur.rowcount
        if not n:
            return f"ℹ️ Skill `{raw}` не найден."
        return f"✅ Skill `{raw}`: enabled={bool(en)}."

    def format_skills_list(self, lang: str = "ru") -> str:
        """Список skills для чата; `lang` — ru|en (заголовки и пустое состояние)."""
        with self._conn() as conn:
            rows = conn.execute(
                """
                SELECT id, name, enabled, length(body) AS blen, updated_at
                FROM omi_skills
                ORDER BY name
                """
            ).fetchall()
        lang = (lang or "ru").lower()[:2]
        if not rows:
            if lang == "en":
                return (
                    "📋 *Skills:* (none). Owner can add via `skill_add` in chat or JSON. "
                    "Optional: drop `*.md` into `omi_skills_extra/` under the workspace."
                )
            return "📋 *Skills:* пусто. Владелец может добавить: `skill_add` или чат. Дополнительно: файлы в `omi_skills_extra/*.md`."

        if lang == "en":
            lines = ["📋 *Registered skills:*\n"]
            sym = "chars"
        else:
            lines = ["📋 *Зарегистрированные skills:*\n"]
            sym = "симв."
        for r in rows:
            on = "✓" if r["enabled"] else "✗"
            lines.append(
                f"  • `#{r['id']}` [{on}] *{r['name']}* — {r['blen']} {sym}, upd. {(r['updated_at'] or '')[:19]}"
            )
        return "\n".join(lines)

    def get_enabled_skills_prompt_block(self, max_chars: int | None = None) -> str:
        """
        Подмешивание в system prompt: БД `omi_skills`, затем **`aims_skills_shared`** (канон с Axi),
        затем **`omi_skills_extra`** — файлы с тем же именем, что уже в shared, **пропускаются**
        (как `axi_skills_extra` у Axi). Уникальные только-Omi правила держите в `omi_skills_extra`
        под именами, которых нет в `aims_skills_shared`, либо в БД.
        """
        if max_chars is None:
            try:
                max_chars = int(os.environ.get("OMI_SKILLS_MAX_CHARS", "24000"))
            except ValueError:
                max_chars = 24000

        parts: list[str] = []
        used = 0

        with self._conn() as conn:
            rows = conn.execute(
                "SELECT name, body FROM omi_skills WHERE enabled = 1 ORDER BY name"
            ).fetchall()

        if rows:
            buf: list[str] = ["[Skills from DB (enabled) — use when relevant:\n"]
            for r in rows:
                chunk = f"### {r['name']}\n{r['body']}\n\n"
                if used + len(chunk) > max_chars:
                    buf.append(f"[DB skills truncated; OMI_SKILLS_MAX_CHARS={max_chars}]\n")
                    break
                buf.append(chunk)
                used += len(chunk)
            buf.append("]")
            parts.append("".join(buf))

        def _budget_left() -> int:
            return max(0, max_chars - sum(len(p) for p in parts))

        shared_root = self.workspace / "aims_skills_shared"
        shared_names: set[str] = set()
        if shared_root.is_dir():
            shared_names = {
                p.name
                for p in shared_root.iterdir()
                if p.is_file() and p.suffix.lower() in (".md", ".txt")
            }

        bl = _budget_left()
        if bl > 400 and shared_root.is_dir():
            files = sorted(shared_root.glob("*.md")) + sorted(shared_root.glob("*.txt"))
            sbuf: list[str] = ["[Shared skills `aims_skills_shared/` (Axi+Omi, canonical files):\n"]
            st = 0
            for fp in files:
                try:
                    raw = fp.read_text(encoding="utf-8", errors="ignore")
                except OSError:
                    continue
                chunk = f"### {fp.name}\n{raw}\n\n"
                if st + len(chunk) > bl:
                    sbuf.append("[Shared skills truncated]\n")
                    break
                sbuf.append(chunk)
                st += len(chunk)
            if len(sbuf) > 1:
                sbuf.append("]")
                parts.append("".join(sbuf))

        bl = _budget_left()
        extra_root = self.workspace / "omi_skills_extra"
        if bl > 400 and extra_root.is_dir():
            files = sorted(extra_root.glob("*.md")) + sorted(extra_root.glob("*.txt"))
            wbuf: list[str] = ["[Omi-only extra skills `omi_skills_extra/` (skip if same filename as shared):\n"]
            wt = 0
            for fp in files:
                if fp.name in shared_names:
                    continue
                try:
                    raw = fp.read_text(encoding="utf-8", errors="ignore")
                except OSError:
                    continue
                chunk = f"### {fp.name}\n{raw}\n\n"
                if wt + len(chunk) > bl:
                    wbuf.append("[Files truncated]\n")
                    break
                wbuf.append(chunk)
                wt += len(chunk)
            if len(wbuf) > 1:
                wbuf.append("]")
                parts.append("".join(wbuf))

        return "\n\n".join(parts) if parts else ""

    def _collect_report_link_candidates(self, name: str) -> list[dict]:
        """Кандидаты для семантического сопоставления: гибридный поиск + расширение по токенам имени."""
        qst = Path(name).stem.replace("_", " ").strip()[:72] or name
        by_id: dict[int, dict] = {}
        try:
            for row in self.search_document_files(qst, limit=22):
                did = row.get("id")
                if did is None:
                    continue
                by_id[int(did)] = dict(row)
        except Exception:
            pass
        tokens = [t for t in re.split(r"[^\w]+", Path(name).stem.lower()) if len(t) > 3][:5]
        if not tokens:
            tokens = [Path(name).stem.lower()[:16]] if Path(name).stem else ["x"]
        if len(by_id) >= 20:
            return list(by_id.values())[:28]
        where_parts: list[str] = []
        params: list[str] = []
        for t in tokens:
            where_parts.append(
                "(lower(file_name) LIKE ? OR lower(COALESCE(original_file_name,'')) LIKE ? "
                "OR lower(COALESCE(canonical_file_name,'')) LIKE ? OR lower(COALESCE(title,'')) LIKE ? "
                "OR lower(COALESCE(summary,'')) LIKE ?)"
            )
            p = f"%{t}%"
            params.extend([p, p, p, p, p])
        sql = f"""
            SELECT id, file_name, aims_process,
                   COALESCE(original_file_name,'') AS ofn,
                   COALESCE(canonical_file_name,'') AS cfn,
                   COALESCE(title,'') AS title,
                   substr(COALESCE(summary,''), 1, 240) AS summary_excerpt
            FROM documents
            WHERE {' OR '.join(where_parts)}
            ORDER BY date_added DESC
            LIMIT 45
        """
        with self._conn() as conn:
            for r in conn.execute(sql, params).fetchall():
                did = int(r[0])
                if did in by_id:
                    continue
                by_id[did] = {
                    "id": did,
                    "file_name": r[1],
                    "aims_process": r[2],
                    "original_file_name": str(r[3] or "").strip() or None,
                    "canonical_file_name": str(r[4] or "").strip() or None,
                    "title": str(r[5] or "").strip() or None,
                    "summary_excerpt": str(r[6] or "").strip() or None,
                    "file_path": "",
                }
        return list(by_id.values())[:28]

    def _report_sync_note_heuristic(self, path: Path, name: str, lang: str) -> str:
        """Резерв, если LLM недоступна или OMI_REPORT_LINK_AI=0 — эвристика по нормализации имён."""
        tokens = [t for t in re.split(r"[^\w]+", Path(name).stem.lower()) if len(t) > 3][:5]
        if not tokens:
            tokens = [Path(name).stem.lower()[:16]] if Path(name).stem else ["x"]

        where_parts: list[str] = []
        params: list[str] = []
        for t in tokens:
            where_parts.append(
                "(lower(file_name) LIKE ? OR lower(COALESCE(original_file_name,'')) LIKE ? "
                "OR lower(COALESCE(canonical_file_name,'')) LIKE ?)"
            )
            p = f"%{t}%"
            params.extend([p, p, p])

        sql = f"""
            SELECT id, file_name, aims_process,
                   COALESCE(original_file_name,'') AS ofn,
                   COALESCE(canonical_file_name,'') AS cfn
            FROM documents
            WHERE {' OR '.join(where_parts)}
            ORDER BY date_added DESC
            LIMIT 60
        """
        best_row: tuple | None = None
        best_score = 0.0
        best_label = ""
        with self._conn() as conn:
            rows = conn.execute(sql, params).fetchall()
            for r in rows:
                sc, lbl = score_report_vs_registry_row(
                    name,
                    file_name=str(r[1] or ""),
                    original_file_name=str(r[3] or "") or None,
                    canonical_file_name=str(r[4] or "") or None,
                )
                if sc > best_score:
                    best_score = sc
                    best_label = lbl
                    best_row = r

        intro_ru = (
            "\n\nСверка с БД (эвристика имён): в реестре имена часто анонимизированы; "
            "сопоставление по original_file_name / canonical / нормализованному stem."
        )
        intro_en = (
            "\n\nDB check (name heuristic): registry file_name is often anonymized; "
            "linking uses original/canonical and normalized stems."
        )

        if best_row is not None and best_score >= 0.88:
            proc = best_row[2] or "—"
            db_fn = best_row[1]
            field_ru = {"file_name": "file_name", "original": "original_file_name", "canonical": "canonical_file_name"}.get(
                best_label, best_label
            )
            if lang == "ru":
                return (
                    f"{intro_ru}\n✓ Высокая вероятность связи с записью id={best_row[0]}, процесс {proc}, "
                    f"в БД file_name=`{db_fn}` (лучшее поле: {field_ru}). "
                    "Не утверждай идентичность содержимого без проверки пользователем."
                )
            return (
                f"{intro_en}\n✓ Likely registry link id={best_row[0]}, process {proc}, "
                f"DB file_name=`{db_fn}` (best field: {best_label}). "
                "Do not assert same content without user verification."
            )

        if best_row is not None and best_score >= 0.5:
            proc = best_row[2] or "—"
            db_fn = best_row[1]
            if lang == "ru":
                return (
                    f"{intro_ru}\n⚠ Слабая связь с id={best_row[0]} (`{db_fn}`, процесс {proc}). "
                    "Считай отчёт отдельным артефактом, пока пользователь не подтвердит соответствие."
                )
            return (
                f"{intro_en}\n⚠ Weak possible link id={best_row[0]} (`{db_fn}`, process {proc}). "
                "Treat report as separate until user confirms."
            )

        if lang == "ru":
            return (
                f"{intro_ru}\n⚠ Явной связи с записями documents по имени не найдено. "
                "Файл в report/ может быть черновиком или вне реестра."
            )
        return f"{intro_en}\n⚠ No confident name-based link to documents. File may be draft or outside registry."

    def report_sync_note(self, path: Path, *, lang: str = "en") -> str:
        """Сверка report/ с реестром: приоритет семантики (Ollama), иначе эвристика имён."""
        report_root = Path(os.environ.get("AIMS_REPORT_DIR", str(self.workspace / "report")))
        try:
            path.resolve().relative_to(report_root.resolve())
        except ValueError:
            return ""
        except OSError:
            return ""
        name = path.name

        intro_ru_ai = (
            "\n\nСверка с БД (семантика, локальная LLM): в реестре имена часто анонимизированы; "
            "сопоставление по смыслу отчёта с title/summary и алиасами имён."
        )
        intro_en_ai = (
            "\n\nDB check (semantic, local LLM): registry names are often anonymized; "
            "matching uses report meaning vs title/summary and name aliases."
        )

        ai_on = os.environ.get("OMI_REPORT_LINK_AI", "1").strip().lower() in ("1", "true", "yes", "on")
        if ai_on:
            candidates = self._collect_report_link_candidates(name)
            if candidates:
                excerpt = read_report_excerpt(path)
                ai = semantic_match_report_to_candidates(
                    report_filename=name,
                    excerpt=excerpt,
                    candidates=candidates,
                )
                if ai:
                    out = format_sync_note_from_ai(
                        ai,
                        lang=lang,
                        intro_ru=intro_ru_ai,
                        intro_en=intro_en_ai,
                    )
                    if (out or "").strip():
                        return out
        return self._report_sync_note_heuristic(path, name, lang)

    def fetch_documents_for_bundle(
        self,
        *,
        process_code: str | None = None,
        doc_ids: list[int] | None = None,
        limit: int = 120,
    ) -> list[dict]:
        """Строки для генерации документации."""
        limit = max(1, min(int(limit), 500))
        with self._conn() as conn:
            if doc_ids:
                uniq = [int(x) for x in doc_ids[:200]]
                if not uniq:
                    return []
                ph = ",".join(["?"] * len(uniq))
                sql = f"""
                    SELECT id, file_name, aims_process, aims_element, iso_clause, title, summary,
                           date_added, is_master, is_anonymized, file_path
                    FROM documents
                    WHERE id IN ({ph})
                    ORDER BY aims_process, file_name
                """
                rows = conn.execute(sql, uniq).fetchall()
            elif process_code and str(process_code).strip():
                pc = str(process_code).strip().upper()
                rows = conn.execute(
                    """
                    SELECT id, file_name, aims_process, aims_element, iso_clause, title, summary,
                           date_added, is_master, is_anonymized, file_path
                    FROM documents
                    WHERE aims_process = ?
                    ORDER BY file_name
                    LIMIT ?
                    """,
                    (pc, limit),
                ).fetchall()
            else:
                rows = conn.execute(
                    """
                    SELECT id, file_name, aims_process, aims_element, iso_clause, title, summary,
                           date_added, is_master, is_anonymized, file_path
                    FROM documents
                    ORDER BY date_added DESC
                    LIMIT ?
                    """,
                    (limit,),
                ).fetchall()
        out: list[dict] = []
        for r in rows:
            out.append({
                "id": r["id"],
                "file_name": r["file_name"],
                "aims_process": r["aims_process"],
                "aims_element": r["aims_element"],
                "iso_clause": r["iso_clause"],
                "title": r["title"],
                "summary": r["summary"],
                "date_added": r["date_added"],
                "is_master": bool(r["is_master"]) if r["is_master"] is not None else False,
                "is_anonymized": bool(r["is_anonymized"]) if r["is_anonymized"] is not None else False,
                "file_path": r["file_path"],
            })
        return out

    @staticmethod
    def _compute_file_hash(path: Path) -> str | None:
        """SHA-256 of file bytes for content-based dedup. Returns None on read error."""
        try:
            h = hashlib.sha256()
            with open(path, "rb") as fh:
                for chunk in iter(lambda: fh.read(65536), b""):
                    h.update(chunk)
            return h.hexdigest()
        except Exception:
            return None

    def register_generated_bundle(
        self,
        file_path: Path | str,
        *,
        aims_process: str | None = None,
        master_doc_json: str | None = None,
    ) -> str:
        """
        Регистрирует сгенерированный .docx в таблице `documents` (после одобрения в чате).
        Файл должен лежать внутри `AIMS_WORKSPACE` (в т.ч. `generated/`).
        """
        file_path = Path(file_path).resolve()
        ws = self.workspace.resolve()
        try:
            file_path.relative_to(ws)
        except ValueError:
            return "❌ Файл вне рабочего каталога workspace."
        if not file_path.is_file():
            return "❌ Файл не найден."

        fn = file_path.name
        title = fn.rsplit(".", 1)[0] if "." in fn else fn
        ft = fn.rsplit(".", 1)[-1].lower() if "." in fn else ""
        now = datetime.now().isoformat()
        created_fs, mtime_fs = _file_times_iso(str(file_path))
        summary = f"Сводный пакет документации (сгенерировано Omi). Файл: {fn}"
        ap = (aims_process or "DOCGEN").strip().upper()[:16]
        notes = json.dumps(
            {"source": "docgen_bundle", "generated": True, "registered_at": now},
            ensure_ascii=False,
        )
        fp = str(file_path)
        content_hash = self._compute_file_hash(file_path)

        with self._conn() as conn:
            # Content-hash dedup: if identical content already registered under a
            # different path, skip the INSERT and return the canonical registration.
            if content_hash:
                dup = conn.execute(
                    "SELECT id, file_path FROM documents"
                    " WHERE content_hash = ? AND file_path != ?",
                    (content_hash, fp),
                ).fetchone()
                if dup:
                    dup_name = Path(dup[1]).name
                    return (
                        f"Дубликат: содержимое уже зарегистрировано как `{dup_name}`"
                        f" (hash={content_hash[:12]}…)."
                    )

            row = conn.execute("SELECT id FROM documents WHERE file_path = ?", (fp,)).fetchone()
            if row:
                conn.execute(
                    """
                    UPDATE documents SET
                        file_name = ?, file_type = ?, title = ?, summary = ?, aims_process = ?,
                        is_master = 0, is_anonymized = 1, date_modified = ?,
                        file_created_at = ?, file_mtime_at = ?,
                        original_file_name = COALESCE(original_file_name, file_name),
                        canonical_file_name = COALESCE(canonical_file_name, file_name),
                        source_filename = COALESCE(source_filename, ?),
                        stored_path = COALESCE(stored_path, ?),
                        process_code = COALESCE(process_code, ?),
                        anonymized_result_path = COALESCE(anonymized_result_path, ''),
                        master_doc_json = COALESCE(?, master_doc_json),
                        content_hash = COALESCE(?, content_hash),
                        created_at = COALESCE(created_at, ?),
                        updated_at = ?,
                        source = ?, notes = ?
                    WHERE file_path = ?
                    """,
                    (
                        fn,
                        ft,
                        title,
                        summary,
                        ap,
                        now,
                        created_fs or now,
                        mtime_fs or now,
                        fn,
                        fp,
                        ap,
                        master_doc_json,
                        content_hash,
                        now,
                        now,
                        "docgen_bundle",
                        notes,
                        fp,
                    ),
                )
            else:
                conn.execute(
                    """
                    INSERT INTO documents (
                        file_path, file_name, file_type, title, summary, aims_process,
                        is_master, is_anonymized, language, date_added, date_modified,
                        file_created_at, file_mtime_at, original_file_name, canonical_file_name,
                        source_filename, stored_path, process_code, anonymized_result_path,
                        master_doc_json, content_hash, source, notes, created_at, updated_at
                    ) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
                    """,
                    (
                        fp,
                        fn,
                        ft,
                        title,
                        summary,
                        ap,
                        0,
                        1,
                        "en",
                        now,
                        now,
                        created_fs or now,
                        mtime_fs or now,
                        fn,
                        fn,
                        fn,
                        fp,
                        ap,
                        "",
                        master_doc_json,
                        content_hash,
                        "docgen_bundle",
                        notes,
                        now,
                        now,
                    ),
                )
            conn.commit()
        self._log("register_docgen_bundle", detail=fn[:400])
        return f"Зарегистрировано в реестре: `{fn}` (процесс `{ap}`)."

    # ── Правила поведения от владельца (omi_config) ───────────

    _OWNER_RULES_KEY = "owner_behavior_rules"

    def get_owner_rules_text(self) -> str:
        """Текст правил, подмешивается в системный промпт агента."""
        try:
            with self._conn() as conn:
                row = conn.execute(
                    "SELECT value FROM omi_config WHERE key = ?",
                    (self._OWNER_RULES_KEY,),
                ).fetchone()
            return (row[0] or "").strip() if row else ""
        except Exception:
            return ""

    def append_owner_rule(self, rule: str, actor: str | None = None) -> None:
        """Добавить строку правила (с меткой времени)."""
        rule = (rule or "").strip()
        if not rule:
            return
        now = datetime.now().isoformat()
        line = f"[{now}] {rule}"
        existing = self.get_owner_rules_text()
        new_val = f"{existing}\n{line}".strip() if existing else line
        with self._conn() as conn:
            conn.execute(
                """
                INSERT INTO omi_config(key, value, updated)
                VALUES (?, ?, ?)
                ON CONFLICT(key) DO UPDATE SET
                    value = excluded.value,
                    updated = excluded.updated
                """,
                (self._OWNER_RULES_KEY, new_val, now),
            )
            conn.commit()
        self._log("add_rule", detail=line[:500], actor=actor)

    def format_owner_rules_for_user(self) -> str:
        """Сообщение для list_rules в Telegram."""
        raw = self.get_owner_rules_text()
        if not raw:
            return "📋 *Правил от владельца пока нет.* Добавь текстом: «добавь правило: …»"
        lines = raw.split("\n")
        out = ["📋 *Сохранённые правила поведения:*\n"]
        for i, line in enumerate(lines, 1):
            if line.strip():
                out.append(f"{i}. `{line[:900]}`")
        return "\n".join(out)

    def get_config_value(self, key: str) -> str | None:
        try:
            with self._conn() as conn:
                row = conn.execute(
                    "SELECT value FROM omi_config WHERE key = ?", (key,)
                ).fetchone()
            return str(row[0]) if row and row[0] is not None else None
        except Exception:
            return None

    def set_config_value(self, key: str, value: str) -> None:
        with self._conn() as conn:
            conn.execute(
                """
                INSERT INTO omi_config(key, value, updated)
                VALUES (?, ?, ?)
                ON CONFLICT(key) DO UPDATE SET
                    value = excluded.value,
                    updated = excluded.updated
                """,
                (key, value, datetime.now().isoformat()),
            )
            conn.commit()

    # ── 1. Статус ─────────────────────────────────────────────

    def get_status(self) -> dict:
        with self._conn() as conn:
            total  = conn.execute("SELECT COUNT(*) FROM documents").fetchone()[0]
            master = conn.execute("SELECT COUNT(*) FROM documents WHERE is_master=1").fetchone()[0]
            anon   = conn.execute("SELECT COUNT(*) FROM documents WHERE is_anonymized=1").fetchone()[0]
            procs  = conn.execute("SELECT code, name FROM omi_processes ORDER BY code").fetchall()
            proc_counts = []
            for p in procs:
                c = conn.execute(
                    "SELECT COUNT(*) FROM documents WHERE aims_process=?", (p["code"],)
                ).fetchone()[0]
                proc_counts.append({"code": p["code"], "name": p["name"], "count": c})

        return {
            "db_path":    str(self.db_path),
            "workspace":  str(self.workspace),
            "total_docs": total,
            "master_docs": master,
            "anon_docs":  anon,
            "processes":  proc_counts,
            "platform":   self._detect_platform(),
        }

    def gather_ocr_skipped_context(self, *, max_files: int = 15) -> tuple[str, int, int]:
        """
        Собирает краткое описание файлов в inbox/Skipped и счётчиков .ocr_fail для AI-диагностики.
        Возвращает (текст для промпта, число файлов в Skipped в выборке, всего в папке).
        """
        skipped = self.workspace / "inbox" / "Skipped"
        fail_dir = self.workspace / "inbox" / ".ocr_fail"
        if not skipped.is_dir():
            return ("(папка inbox/Skipped отсутствует)", 0, 0)
        all_files = sorted(
            [p for p in skipped.iterdir() if p.is_file()],
            key=lambda p: p.stat().st_mtime,
            reverse=True,
        )
        total_n = len(all_files)
        take = all_files[: max(1, min(max_files, 30))]
        lines: list[str] = []
        for p in take:
            ext = p.suffix.lower()
            sz_kb = max(1, p.stat().st_size // 1024)
            cnt_raw = ""
            if fail_dir.is_dir():
                cf = fail_dir / f"{p.stem}.cnt"
                if cf.is_file():
                    try:
                        cnt_raw = (cf.read_text(encoding="utf-8", errors="ignore") or "").strip()[:120]
                    except OSError:
                        cnt_raw = "?"
            lines.append(f"- name={p.name} ext={ext} size_kb≈{sz_kb} ocr_fail_cnt_file={cnt_raw or 'none'}")
        body = f"Skipped folder: {total_n} file(s), sample:\n" + "\n".join(lines)
        return (body, len(take), total_n)

    def _safe_inbox_filename(self, name: str) -> str | None:
        """Только basename без path traversal; разрешены типичные символы имён файлов."""
        if not isinstance(name, str):
            return None
        base = Path(name.strip()).name
        if not base or len(base) > 255:
            return None
        if base != name.strip():
            return None
        if ".." in base:
            return None
        return base

    def list_skipped_inbox_files(self, *, limit: int = 40) -> str:
        """Список имён в inbox/Skipped для выбора точного filename в skipped_ops."""
        skipped = self.workspace / "inbox" / "Skipped"
        if not skipped.is_dir():
            return "📂 `inbox/Skipped` отсутствует."
        files = sorted(
            [p for p in skipped.iterdir() if p.is_file()],
            key=lambda p: p.stat().st_mtime,
            reverse=True,
        )[: max(1, min(limit, 80))]
        if not files:
            return "✅ В `inbox/Skipped` нет файлов."
        lines = [f"📋 `inbox/Skipped` — {len(files)} (последние по времени):"]
        for p in files:
            sz = max(1, p.stat().st_size // 1024)
            lines.append(f"  • `{p.name}` ({sz} KB)")
        return "\n".join(lines)

    def skipped_ops(
        self,
        op: str,
        filename: str,
        *,
        dest: str | None = None,
        context_text: str | None = None,
        is_owner: bool = False,
    ) -> str:
        """
        Операции над файлами в inbox/Skipped — исполнение советов диагностики OCR.

        op: delete | requeue | renormalize | clear_fail | auto
        dest (requeue): inbox | income — куда переложить для повторной обработки (ocr-watcher читает корень inbox).
        """
        op = (op or "").strip().lower()
        op_compact = re.sub(r"[^a-z]", "", op)
        if op_compact in ("requeuetoinbox", "requeueinbox"):
            op = "requeue"
            if not dest:
                dest = "inbox"
        elif op_compact in ("requeuetoincome", "requeueincome"):
            op = "requeue"
            if not dest:
                dest = "income"
        elif op_compact in ("clearfail", "clearretry", "resetfail"):
            op = "clear_fail"
        elif op_compact in ("renormalize", "normalizename", "renamewithunderscores"):
            op = "renormalize"
        base = self._safe_inbox_filename(filename)
        if not base:
            return "❌ Некорректное имя файла (только имя, без путей)."
        # Allow AI/context-driven execution without explicit op.
        if op in ("", "auto", "context", "ai"):
            t = (context_text or "").lower()
            if base.startswith("~$"):
                op = "delete"
            elif any(k in t for k in ("удал", "delete", "drop", "trash")):
                op = "delete"
            elif any(k in t for k in ("переимен", "renormalize", "normalize", "пробел", "spaces")):
                op = "renormalize"
            elif any(k in t for k in ("clear_fail", "сброс", "reset fail", "retry counter")):
                op = "clear_fail"
            elif any(
                k in t
                for k in (
                    "повтор",
                    "requeue",
                    "retry",
                    "заново",
                    "в очередь",
                    "queue",
                    "перелож",
                    "верни в inbox",
                )
            ):
                op = "requeue"
            elif " " in base:
                op = "renormalize"
            else:
                op = "requeue"
            if op == "requeue" and (dest is None or not str(dest).strip()):
                if any(k in t for k in ("income", "inbox/income")):
                    dest = "income"
                else:
                    dest = "inbox"
        skipped_dir = self.workspace / "inbox" / "Skipped"
        fail_dir = self.workspace / "inbox" / ".ocr_fail"
        src = skipped_dir / base

        if op == "delete":
            if not src.is_file():
                return f"❌ Нет файла `Skipped/{base}`."
            is_lock = base.startswith("~$")
            if not is_owner and not is_lock:
                return (
                    "⚠️ Удаление доступно только *владельцу* (`OMI_OWNER_CHAT_IDS`), "
                    "кроме lock-файлов Word (`~$…`)."
                )
            try:
                src.unlink()
                self._log("skipped_delete", detail=base[:200])
                return f"✅ Удалён `inbox/Skipped/{base}`."
            except OSError as e:
                return f"❌ Не удалось удалить: {e}"

        if op == "requeue":
            if not src.is_file():
                return f"❌ Нет файла `Skipped/{base}`."
            d = (dest or "inbox").strip().lower()
            if d not in ("inbox", "income"):
                d = "inbox"
            if d == "income":
                target_dir = self.workspace / "inbox" / "income"
            else:
                target_dir = self.workspace / "inbox"
            target_dir.mkdir(parents=True, exist_ok=True)
            dst = target_dir / base
            if dst.exists():
                return (
                    f"❌ Уже существует `{dst.name}` в целевой папке — устраните конфликт имён вручную."
                )
            try:
                shutil.move(str(src), str(dst))
                self._log(
                    "skipped_requeue",
                    detail=json.dumps({"from": str(src), "to": str(dst)}, ensure_ascii=False)[:500],
                )
                rel = f"inbox/income/{base}" if d == "income" else f"inbox/{base}"
                return f"✅ Перенесён → `{rel}` (ожидает обработки ocr-watcher / пайплайна)."
            except OSError as e:
                return f"❌ Ошибка переноса: {e}"

        if op == "renormalize":
            if not src.is_file():
                return f"❌ Нет файла `Skipped/{base}`."
            if " " not in base:
                return f"ℹ️ В имени `{base}` нет пробелов — переименование не требуется."
            stem, suf = Path(base).stem, Path(base).suffix
            new_base = stem.replace(" ", "_") + suf
            if new_base == base:
                return "ℹ️ Имя не изменилось."
            dst = skipped_dir / new_base
            if dst.exists():
                return f"❌ Цель `{new_base}` уже существует."
            try:
                shutil.move(str(src), str(dst))
                self._log("skipped_renormalize", detail=f"{base} -> {new_base}")
                return f"✅ Переименовано в `Skipped/{new_base}` (пробелы → _)."
            except OSError as e:
                return f"❌ Ошибка: {e}"

        if op == "clear_fail":
            stem = Path(base).stem
            cnt = fail_dir / f"{stem}.cnt"
            if not cnt.is_file():
                return f"ℹ️ Счётчика `.ocr_fail/{stem}.cnt` нет (или уже сброшен)."
            try:
                cnt.unlink()
                self._log("skipped_clear_fail", detail=stem[:120])
                return f"✅ Сброшен счётчик повторов OCR: `.ocr_fail/{stem}.cnt`."
            except OSError as e:
                return f"❌ Не удалось удалить счётчик: {e}"

        return (
            f"❌ Неизвестная операция skipped_ops: `{op}`. "
            "Допустимо: delete, requeue, renormalize, clear_fail, auto."
        )

    def skipped_compare(
        self,
        *,
        apply_requeue_non_duplicates: bool = False,
        dest: str = "inbox",
        limit: int = 200,
    ) -> str:
        """
        Сравнение файлов в inbox/Skipped по размеру и sha256.
        Опционально: перенести уникальные (не-дубликаты) в inbox/income для повторной обработки.
        """
        skipped_dir = self.workspace / "inbox" / "Skipped"
        if not skipped_dir.is_dir():
            return "📂 `inbox/Skipped` отсутствует."

        files = sorted(
            [p for p in skipped_dir.iterdir() if p.is_file()],
            key=lambda p: p.stat().st_mtime,
            reverse=True,
        )[: max(1, min(int(limit or 200), 500))]
        if not files:
            return "✅ В `inbox/Skipped` нет файлов."

        def _sha256(path: Path) -> str:
            h = hashlib.sha256()
            with path.open("rb") as f:
                while True:
                    chunk = f.read(1024 * 1024)
                    if not chunk:
                        break
                    h.update(chunk)
            return h.hexdigest()

        by_sig: dict[tuple[int, str], list[Path]] = {}
        errors: list[str] = []
        for p in files:
            try:
                sz = int(p.stat().st_size)
                sig = (sz, _sha256(p))
                by_sig.setdefault(sig, []).append(p)
            except OSError as e:
                errors.append(f"`{p.name}`: {str(e)[:120]}")

        duplicate_groups = [grp for grp in by_sig.values() if len(grp) > 1]
        duplicate_names: set[str] = set()
        for grp in duplicate_groups:
            for p in grp:
                duplicate_names.add(p.name)

        target = (dest or "inbox").strip().lower()
        if target not in ("inbox", "income"):
            target = "inbox"
        target_dir = self.workspace / "inbox" / ("income" if target == "income" else "")
        moved: list[str] = []
        conflicts: list[str] = []
        move_errors: list[str] = []
        if apply_requeue_non_duplicates:
            target_dir.mkdir(parents=True, exist_ok=True)
            for p in files:
                if p.name in duplicate_names:
                    continue
                dst = target_dir / p.name
                if dst.exists():
                    conflicts.append(p.name)
                    continue
                try:
                    shutil.move(str(p), str(dst))
                    moved.append(p.name)
                except OSError as e:
                    move_errors.append(f"`{p.name}`: {str(e)[:120]}")

        lines = [
            f"📊 Сравнение `inbox/Skipped`: файлов={len(files)}, duplicate_groups={len(duplicate_groups)}, duplicates_total={len(duplicate_names)}",
        ]
        if duplicate_groups:
            lines.append("\n🔁 Группы полных дублей (одинаковые размер+sha256):")
            for i, grp in enumerate(duplicate_groups[:30], 1):
                size_kb = max(1, grp[0].stat().st_size // 1024)
                names = ", ".join(f"`{p.name}`" for p in grp[:8])
                more = "" if len(grp) <= 8 else f" ... +{len(grp) - 8}"
                lines.append(f"  {i}. size={size_kb} KB: {names}{more}")

        unique_n = len(files) - len(duplicate_names)
        lines.append(f"\n✅ Уникальных по содержимому: {unique_n}")

        if apply_requeue_non_duplicates:
            rel_target = "inbox/income" if target == "income" else "inbox"
            lines.append(f"🚚 Перенос уникальных в `{rel_target}`: {len(moved)}")
            if moved:
                lines.extend([f"  • `{n}`" for n in moved[:60]])
            if conflicts:
                lines.append(f"⚠️ Конфликты имён (уже есть в целевой папке): {len(conflicts)}")
                lines.extend([f"  • `{n}`" for n in conflicts[:30]])
            if move_errors:
                lines.append(f"❌ Ошибки переноса: {len(move_errors)}")
                lines.extend([f"  • {t}" for t in move_errors[:20]])

        if errors:
            lines.append(f"\n⚠️ Ошибки чтения файлов: {len(errors)}")
            lines.extend([f"  • {e}" for e in errors[:20]])

        self._log(
            "skipped_compare",
            detail=json.dumps(
                {
                    "files": len(files),
                    "duplicate_groups": len(duplicate_groups),
                    "duplicates_total": len(duplicate_names),
                    "moved": len(moved),
                    "target": target,
                    "apply": bool(apply_requeue_non_duplicates),
                },
                ensure_ascii=False,
            )[:500],
        )
        return "\n".join(lines)

    def list_processes(self) -> list:
        with self._conn() as conn:
            rows = conn.execute(
                """
                SELECT
                    p.code,
                    p.name,
                    COALESCE(COUNT(d.id), 0) AS count
                FROM omi_processes p
                LEFT JOIN documents d
                    ON d.aims_process = p.code
                GROUP BY p.code, p.name
                ORDER BY p.code
                """
            ).fetchall()
        return [
            {"code": r["code"], "name": r["name"], "count": int(r["count"] or 0)}
            for r in rows
        ]

    # ── 2. Поиск ──────────────────────────────────────────────

    def search_documents(
        self, query: str, process: str = None, limit: int = 10, lang: str = "en"
    ) -> str:
        lang = "ru" if (lang or "").lower().startswith("ru") else "en"
        # Split into words for word-by-word matching (handles underscore/space mismatches)
        words = [w.strip() for w in query.split() if len(w.strip()) > 1]
        if not words:
            words = [query]
        per_word_clause = "(keywords LIKE ? OR title LIKE ? OR summary LIKE ? OR file_name LIKE ?)"
        where_clause = " AND ".join([per_word_clause] * len(words))
        sql = f"""
            SELECT id, file_name, aims_process, title, date_added, file_path,
                   COALESCE(file_created_at, '') AS fca,
                   COALESCE(file_mtime_at, '') AS fma
            FROM documents
            WHERE ({where_clause})
        """
        params: list = []
        for w in words:
            wl = f"%{w}%"
            params.extend([wl, wl, wl, wl])
        if process:
            sql += " AND aims_process = ?"
            params.append(process)
        sql += f" ORDER BY date_added DESC LIMIT {limit}"

        with self._conn() as conn:
            rows = conn.execute(sql, params).fetchall()

        if not rows:
            if lang == "ru":
                return f"🔍 По запросу *«{query}»* ничего не найдено."
            return f"🔍 No results for *«{query}»*."

        if lang == "ru":
            lines = [f"🔍 *Найдено: {len(rows)} документов*\n"]
            reg_lbl = "внесено в реестр"
            fs_lbl = "файл (mtime)"
        else:
            lines = [f"🔍 *Found: {len(rows)} document(s)*\n"]
            reg_lbl = "registered"
            fs_lbl = "file mtime"
        for r in rows:
            proc = r["aims_process"] or "—"
            reg = str(r["date_added"] or "").replace("T", " ")[:19] or "—"
            fs = str(r["fma"] or "").replace("T", " ")[:19] or _file_mtime_short(str(r["file_path"] or "")) or "—"
            lines.append(
                f"`[{r['id']}]` `{proc}` {r['file_name']}\n"
                f"       📥 {reg_lbl}: {reg} · 💾 {fs_lbl}: {fs}"
            )
        return "\n".join(lines)

    def search_document_files(self, query: str, limit: int = 10) -> list:
        """Return rows matching query, ordered by Reciprocal Rank Fusion (RRF) of FTS5, LIKE, and Qdrant.

        RRF formula: score = sum(1 / (k + rank_in_source)) across all retrieval sources.
        k=60 (standard). Rows retain _fts_rank and _qdrant_score for downstream soft scoring.
        """
        _RRF_K = 60
        fetch = max(limit * 3, 30)
        words = [w for w in query.split() if len(w) > 2]

        rrf: dict[int, float] = {}   # doc_id -> accumulated RRF score
        meta: dict[int, dict] = {}   # doc_id -> row fields

        def _register(r, source: str) -> int:
            did = r["id"]
            ofn = None
            cfn = None
            if "original_file_name" in r.keys():
                t = (r["original_file_name"] or "").strip()
                ofn = t or None
            if "canonical_file_name" in r.keys():
                t = (r["canonical_file_name"] or "").strip()
                cfn = t or None
            tit = ((r["title"] or "").strip()[:500]) if "title" in r.keys() else None
            sex = None
            if "summary_excerpt" in r.keys():
                sex = ((r["summary_excerpt"] or "").strip()[:320]) or None
            if did not in meta:
                meta[did] = {
                    "id": did,
                    "file_name": r["file_name"],
                    "file_path": r["file_path"],
                    "aims_process": r["aims_process"],
                    "date_added": r["date_added"] if "date_added" in r.keys() else None,
                    "original_file_name": ofn,
                    "canonical_file_name": cfn,
                    "title": tit,
                    "summary_excerpt": sex,
                    "_source": source,
                }
            else:
                if ofn and not meta[did].get("original_file_name"):
                    meta[did]["original_file_name"] = ofn
                if cfn and not meta[did].get("canonical_file_name"):
                    meta[did]["canonical_file_name"] = cfn
                if tit and not meta[did].get("title"):
                    meta[did]["title"] = tit
                if sex and not meta[did].get("summary_excerpt"):
                    meta[did]["summary_excerpt"] = sex
            return did

        with self._conn() as conn:
            # 1. FTS5 — ordered by BM25 ascending (lower value = more relevant)
            try:
                fts_q = " OR ".join(words) if words else query
                fts_rows = conn.execute(
                    "SELECT d.id, d.file_name, d.file_path, d.aims_process, d.date_added,"
                    " d.original_file_name, d.canonical_file_name, d.title,"
                    " substr(COALESCE(d.summary, ''), 1, 240) AS summary_excerpt,"
                    " bm25(documents_fts) AS fts_rank "
                    "FROM documents d JOIN documents_fts f ON d.id = f.rowid "
                    "WHERE documents_fts MATCH ? ORDER BY fts_rank ASC LIMIT ?",
                    [fts_q, fetch],
                ).fetchall()
                for rank, r in enumerate(fts_rows):
                    did = _register(r, "fts")
                    meta[did]["_fts_rank"] = r["fts_rank"]
                    rrf[did] = rrf.get(did, 0.0) + 1.0 / (_RRF_K + rank)
            except Exception:
                pass

            # 2. LIKE per word — ordered by date_added DESC; keep best rank per doc
            #    Escape _ and % in LIKE patterns (they are wildcards in SQLite)
            like_best: dict[int, int] = {}
            _like_sql = (
                "SELECT id, file_name, file_path, aims_process, date_added,"
                " original_file_name, canonical_file_name, title,"
                " substr(COALESCE(summary, ''), 1, 240) AS summary_excerpt FROM documents "
                "WHERE (file_name LIKE ? ESCAPE '\\' OR source_filename LIKE ? ESCAPE '\\'"
                "       OR title LIKE ? ESCAPE '\\' OR summary LIKE ? ESCAPE '\\'"
                "       OR keywords LIKE ? ESCAPE '\\'"
                "       OR COALESCE(original_file_name,'') LIKE ? ESCAPE '\\'"
                "       OR COALESCE(canonical_file_name,'') LIKE ? ESCAPE '\\') "
                "ORDER BY date_added DESC LIMIT ?"
            )
            for word in (words or [query]):
                wl_escaped = word.replace("\\", "\\\\").replace("%", "\\%").replace("_", "\\_")
                wl = f"%{wl_escaped}%"
                like_rows = conn.execute(
                    _like_sql, [wl, wl, wl, wl, wl, wl, wl, fetch],
                ).fetchall()
                for rank, r in enumerate(like_rows):
                    did = _register(r, "like")
                    if did not in like_best or rank < like_best[did]:
                        like_best[did] = rank
            for did, rank in like_best.items():
                rrf[did] = rrf.get(did, 0.0) + 1.0 / (_RRF_K + rank)

            # 2b. Direct ID search — if query contains a number, try matching
            #     file_name patterns like "1_NNNNN" or doc id directly
            _nums = re.findall(r"\d{3,}", query)
            for _num in _nums:
                _id_rows = conn.execute(
                    "SELECT id, file_name, file_path, aims_process, date_added,"
                    " original_file_name, canonical_file_name, title,"
                    " substr(COALESCE(summary, ''), 1, 240) AS summary_excerpt"
                    " FROM documents"
                    " WHERE id = ? OR file_name LIKE ? ESCAPE '\\'",
                    [int(_num), f"%\\_{_num}%"],
                ).fetchall()
                for rank, r in enumerate(_id_rows):
                    did = _register(r, "id_match")
                    if did not in like_best or 0 < like_best[did]:
                        like_best[did] = 0
                        rrf[did] = rrf.get(did, 0.0) + 1.0 / (_RRF_K + 0)

        # 3a. Qdrant semantic search (external, optional)
        _qdrant_ok = False
        try:
            import omi_qdrant
            q_hits = omi_qdrant.search(query, limit=fetch)
            if q_hits:
                _qdrant_ok = True
                missing_ids = [h["doc_id"] for h in q_hits if h["doc_id"] not in meta]
                if missing_ids:
                    placeholders = ",".join("?" * len(missing_ids))
                    with self._conn() as conn:
                        for r in conn.execute(
                            f"SELECT id, file_name, file_path, aims_process, date_added,"
                            f" original_file_name, canonical_file_name, title,"
                            f" substr(COALESCE(summary, ''), 1, 240) AS summary_excerpt"
                            f" FROM documents WHERE id IN ({placeholders})",
                            missing_ids,
                        ).fetchall():
                            _register(r, "qdrant")
                for rank, h in enumerate(q_hits):
                    did = h["doc_id"]
                    if did in meta:
                        meta[did]["_qdrant_score"] = h["score"]
                        rrf[did] = rrf.get(did, 0.0) + 1.0 / (_RRF_K + rank)
        except Exception:
            pass

        # 3b. Local SQLite vector search (cosine similarity) — fallback when Qdrant unavailable
        if not _qdrant_ok:
            try:
                from aims_search import embed_text as _embed_text, _blob_to_vec, _cosine
                q_vec = _embed_text(query[:8000])
                if q_vec:
                    with self._conn() as conn:
                        m_name = __import__("os").environ.get("OMI_EMBED_MODEL", "nomic-embed-text").strip() or "nomic-embed-text"
                        vec_rows = conn.execute(
                            "SELECT doc_id, vector FROM document_embeddings WHERE model = ?",
                            (m_name,),
                        ).fetchall()
                    vec_scored: list[tuple[int, float]] = []
                    for vid, blob in vec_rows:
                        try:
                            score = _cosine(q_vec, _blob_to_vec(blob))
                            if score > 0.1:
                                vec_scored.append((vid, score))
                        except Exception:
                            pass
                    vec_scored.sort(key=lambda x: x[1], reverse=True)
                    missing_vec = [vid for vid, _ in vec_scored[:fetch] if vid not in meta]
                    if missing_vec:
                        placeholders = ",".join("?" * len(missing_vec))
                        with self._conn() as conn:
                            for r in conn.execute(
                                f"SELECT id, file_name, file_path, aims_process, date_added,"
                                f" original_file_name, canonical_file_name, title,"
                                f" substr(COALESCE(summary, ''), 1, 240) AS summary_excerpt"
                                f" FROM documents WHERE id IN ({placeholders})",
                                missing_vec,
                            ).fetchall():
                                _register(r, "vector")
                    for rank, (vid, vscore) in enumerate(vec_scored[:fetch]):
                        if vid in meta:
                            meta[vid]["_qdrant_score"] = vscore
                            rrf[vid] = rrf.get(vid, 0.0) + 1.0 / (_RRF_K + rank)
            except Exception:
                pass

        # Sort by RRF score descending and attach score to each row
        results = []
        for did, score in sorted(rrf.items(), key=lambda x: x[1], reverse=True):
            if did in meta:
                row = dict(meta[did])
                row["_rrf_score"] = round(score, 6)
                results.append(row)

        return results[:limit]

    def hybrid_search_documents(
        self,
        query: str,
        *,
        process: str | None = None,
        is_master: bool | None = None,
        year: str | None = None,
        doc_type: str | None = None,
        limit: int = 10,
    ) -> list[dict]:
        """Soft-score rerank over RRF-merged results (FTS5 + LIKE + Qdrant) with structural filters.

        RRF base (from search_document_files) captures cross-source relevance rank fusion.
        Structural filters (process, year, doc_type, master) apply signed bonuses/penalties on top.
        """
        rows = self.search_document_files(query, limit=max(10, limit * 4))
        tokens = [t for t in re.split(r"[\s_\-]+", query.lower()) if len(t) > 1]
        now_ts = datetime.now().timestamp()
        scored: list[dict] = []

        for row in rows:
            fn = (row.get("file_name") or "").lower()
            fp = (row.get("file_path") or "").lower()
            proc = (row.get("aims_process") or "").upper()
            reasons: list[str] = []

            # RRF base: normalize _rrf_score (0..~0.05) to 0-3 scale via *60
            # A document ranked #1 in all three sources gets max ~3.0
            rrf_score = row.get("_rrf_score", 0.0) or 0.0
            score = min(3.0, rrf_score * 60.0)
            if rrf_score > 0:
                # Label the dominant source
                src = row.get("_source", "")
                qdrant_score = row.get("_qdrant_score", 0.0) or 0.0
                fts_rank = row.get("_fts_rank")
                if isinstance(fts_rank, (int, float)) and qdrant_score >= 0.65:
                    reasons.append("fts+semantic")
                elif isinstance(fts_rank, (int, float)):
                    reasons.append("fts")
                elif qdrant_score >= 0.65:
                    reasons.append("semantic")
                else:
                    reasons.append("lexical")

            # Lexical token coverage in filename (bonus on top of RRF)
            token_hits = sum(1 for t in tokens if t in fn)
            if tokens and token_hits:
                score += min(1.5, token_hits / max(len(tokens), 1) * 1.5)
                if "keyword match" not in reasons:
                    reasons.append("keyword match")

            # Structural filter bonuses / penalties
            if process:
                if proc == process.upper():
                    score += 2.0
                    reasons.append(f"process {proc}")
                else:
                    score -= 2.0

            if is_master is True:
                if "/master/" in fp.replace("\\", "/") or "master" in fn:
                    score += 1.0
                    reasons.append("master")
                else:
                    score -= 1.0

            if year and year in fn:
                score += 0.8
                reasons.append(f"year {year}")

            if doc_type and doc_type.lower() in fn:
                score += 1.2
                reasons.append(f"type {doc_type}")

            # Recency bonus
            rec_ts = _parse_db_datetime_ts(str(row.get("date_added") or ""))
            if rec_ts:
                days = max(0.0, (now_ts - rec_ts) / 86400.0)
                if days <= 7:
                    score += 0.5
                elif days <= 30:
                    score += 0.2

            r2 = dict(row)
            r2["score"] = round(score, 3)
            r2["reasons"] = reasons[:4]
            scored.append(r2)

        scored.sort(key=lambda x: x.get("score", 0.0), reverse=True)
        return scored[:limit]

    def _contextual_name_for_row(self, row: sqlite3.Row) -> str:
        fn = str(row["file_name"] or "").strip()
        ext = (fn.rsplit(".", 1)[-1].lower() if "." in fn else "docx")[:8]
        title = str(row["title"] or "").strip()
        summary = str(row["summary"] or "").strip()
        proc = str(row["aims_process"] or "").strip()
        stem_src = title or summary or fn.rsplit(".", 1)[0]
        stem = _sanitize_filename_stem(stem_src)
        if proc and not stem.lower().startswith(proc.lower()):
            stem = f"{proc}_{stem}"
        return f"{stem}.{ext}"

    def _extract_text_sample(self, file_path: str, *, max_chars: int = 4000) -> str:
        p = Path(file_path or "")
        if not p.is_file():
            return ""
        ext = p.suffix.lower()
        try:
            if ext in (".txt", ".md", ".log", ".csv"):
                return p.read_text(encoding="utf-8", errors="ignore")[:max_chars]
            if ext in (".docx", ".doc"):
                try:
                    from docx import Document  # type: ignore
                    doc = Document(str(p))
                    text = "\n".join(par.text for par in doc.paragraphs if par.text)
                    if text.strip():
                        return text[:max_chars]
                except Exception:
                    pass
                # Fallback for .doc: convert via soffice to txt
                if ext == ".doc":
                    try:
                        import subprocess, tempfile
                        with tempfile.TemporaryDirectory() as tmpdir:
                            r = subprocess.run(
                                ["soffice", "--headless", "--convert-to", "txt:Text",
                                 "--outdir", tmpdir, str(p)],
                                capture_output=True, timeout=60,
                            )
                            out = Path(tmpdir) / (p.stem + ".txt")
                            if out.exists():
                                return out.read_text(encoding="utf-8", errors="ignore")[:max_chars]
                    except Exception:
                        pass
                return ""
            if ext == ".pdf":
                try:
                    import pdfplumber  # type: ignore
                    with pdfplumber.open(str(p)) as pdf:
                        text = "\n".join(pg.extract_text() or "" for pg in pdf.pages[:5])
                    if text.strip():
                        return text[:max_chars]
                except Exception:
                    pass
                return ""
        except Exception:
            return ""
        return ""

    def _ai_rename_suggest_qwen(self, row: sqlite3.Row, sample: str, *, timeout_sec: float = 45.0) -> tuple[str, str] | None:
        try:
            from omi_ollama import ollama_chat  # type: ignore
        except Exception:
            return None
        fn = str(row["file_name"] or "").strip()
        ext = (fn.rsplit(".", 1)[-1].lower() if "." in fn else "docx")[:8]
        proc = str(row["aims_process"] or "").strip()
        title = str(row["title"] or "").strip()
        summary = str(row["summary"] or "").strip()
        model = os.environ.get("OMI_RENAME_MODEL", os.environ.get("OMI_MODEL", "axi_omi_sphere")).strip() or "axi_omi_sphere"
        from ollama_resolve import effective_ollama_base_url

        base_url = effective_ollama_base_url()
        prompt = (
            "Suggest best archival file name based on content.\n"
            "Return ONLY JSON: {\"stem\":\"...\",\"reason\":\"...\"}.\n"
            "Rules: no extension in stem; concise; include process prefix if known; "
            "ASCII letters/digits/_/- only; no dates unless clearly in content.\n\n"
            f"process={proc}\nold_file={fn}\ntitle={title}\nsummary={summary}\n"
            f"content_sample={sample[:3500]}"
        )
        try:
            raw = ollama_chat(
                [{"role": "user", "content": prompt}],
                model=model,
                base_url=base_url,
                timeout=float(timeout_sec),
            )
            s = (raw or "").strip()
            first = s.find("{")
            last = s.rfind("}")
            if first == -1 or last <= first:
                return None
            data = json.loads(s[first:last + 1])
            stem = _sanitize_filename_stem(str(data.get("stem") or ""))
            if not stem:
                return None
            reason = str(data.get("reason") or "qwen")
            if proc and not stem.lower().startswith(proc.lower()):
                stem = f"{proc}_{stem}"
            return f"{stem}.{ext}", reason[:240]
        except Exception:
            return None

    def _ai_rename_suggest_external(self, row: sqlite3.Row, sample: str, *, timeout_sec: float = 25.0) -> tuple[str, str] | None:
        url = os.environ.get("OMI_RENAME_EXTERNAL_URL", "").strip()
        if not url:
            base = os.environ.get("AXI_API_URL", "http://axi-api:8766").rstrip("/")
            if base:
                url = f"{base}/rename/suggest"
        if not url:
            return None
        fn = str(row["file_name"] or "").strip()
        ext = (fn.rsplit(".", 1)[-1].lower() if "." in fn else "docx")[:8]
        payload = {
            "file_name": fn,
            "title": str(row["title"] or ""),
            "summary": str(row["summary"] or ""),
            "aims_process": str(row["aims_process"] or ""),
            "content_sample": sample[:3500],
            "task": "suggest_archive_name",
        }
        token = os.environ.get("OMI_RENAME_EXTERNAL_TOKEN", "").strip()
        headers = {"Content-Type": "application/json"}
        if token:
            headers["Authorization"] = f"Bearer {token}"
        try:
            body = json.dumps(payload, ensure_ascii=False).encode("utf-8")
            req = _ureq.Request(url, data=body, headers=headers, method="POST")
            with _ureq.urlopen(req, timeout=timeout_sec) as resp:
                text = resp.read().decode("utf-8", errors="ignore")
            data = json.loads(text)
            stem = _sanitize_filename_stem(str(data.get("stem") or data.get("name") or ""))
            if not stem:
                return None
            reason = str(data.get("reason") or "external")
            proc = str(row["aims_process"] or "").strip()
            if proc and not stem.lower().startswith(proc.lower()):
                stem = f"{proc}_{stem}"
            return f"{stem}.{ext}", reason[:240]
        except Exception:
            return None

    def suggest_upload_name_from_content(
        self,
        file_path: str,
        original_name: str,
        *,
        use_external_llm: bool = False,
    ) -> tuple[str, str]:
        """
        Suggest a new filename for freshly uploaded file before registration.
        Returns (new_name, reason). Falls back to original_name safely.
        """
        p = Path(file_path or "")
        orig = (original_name or p.name or "document").strip()
        if "." in orig:
            ext = orig.rsplit(".", 1)[-1].lower()[:8]
        else:
            ext = (p.suffix.lower().lstrip(".") if p.suffix else "txt")
        sample_chars = int(os.environ.get("OMI_UPLOAD_RENAME_SAMPLE_CHARS", "5000"))
        sample = self._extract_text_sample(str(p), max_chars=max(1000, sample_chars))
        if not sample:
            return orig, "no_text_sample"

        fake_row = {
            "file_name": orig,
            "title": "",
            "summary": "",
            "aims_process": "",
        }
        suggested = None
        if use_external_llm:
            suggested = self._ai_rename_suggest_external(fake_row, sample)
        if suggested is None:
            suggested = self._ai_rename_suggest_qwen(fake_row, sample)
        if suggested is None and not use_external_llm and _env_flag("OMI_UPLOAD_RENAME_EXTERNAL_FALLBACK", False):
            suggested = self._ai_rename_suggest_external(fake_row, sample)
        if suggested is None:
            return orig, "ai_unavailable"
        name, reason = suggested
        # Keep original extension for upload staging consistency.
        if "." in name:
            stem = name.rsplit(".", 1)[0]
        else:
            stem = name
        new_name = f"{_sanitize_filename_stem(stem)}.{ext}"
        if not new_name or new_name.startswith("."):
            return orig, "invalid_ai_name"
        return new_name, reason

    def delete_document(self, doc_id: int) -> dict:
        """Удаляет запись о документе из реестра по ID. Файл на диске не трогает."""
        with self._conn() as conn:
            row = conn.execute(
                "SELECT id, COALESCE(canonical_file_name, original_file_name, source_filename) as name FROM documents WHERE id = ?",
                (doc_id,)
            ).fetchone()
            if not row:
                return {"deleted": False, "error": f"Document id={doc_id} not found"}
            conn.execute("DELETE FROM documents WHERE id = ?", (doc_id,))
            conn.commit()
        name = row["name"] or f"id={doc_id}"
        self._log("document_delete", detail=f"id={doc_id} name={name}")
        # Remove vectors from Qdrant
        try:
            import omi_qdrant
            omi_qdrant.delete_document(doc_id)
        except Exception:
            pass
        return {"deleted": True, "id": doc_id, "name": name}

    def rename_documents_by_context(
        self,
        file_names: list[str],
        *,
        lang: str = "en",
        apply_changes: bool = True,
        ai_mode: bool = False,
        use_external_llm: bool = False,
    ) -> str:
        """
        Для списка имён подбирает более предметное имя по title/summary/context и пишет его в реестр.
        Сохраняет старое имя в original_file_name, а новое в canonical_file_name + file_name.
        """
        lang = "ru" if (lang or "").lower().startswith("ru") else "en"
        names = [str(n or "").strip() for n in (file_names or []) if str(n or "").strip()]
        if not names:
            return "⚠️ No file names provided." if lang == "en" else "⚠️ Не переданы имена файлов."

        done: list[str] = []
        miss: list[str] = []
        fail: list[str] = []
        deleted: list[str] = []
        dupes: list[str] = []
        with self._conn() as conn:
            for fn in names:
                row = conn.execute(
                    """
                    SELECT id, file_name, file_path, title, summary, aims_process,
                           COALESCE(original_file_name, '') AS old_name
                    FROM documents
                    WHERE file_name = ?
                    ORDER BY COALESCE(date_added, date_modified, '') DESC
                    LIMIT 1
                    """,
                    (fn,),
                ).fetchone()
                if not row:
                    # Fallback: file not in documents DB — search filesystem in known dirs
                    _fallback_path: Path | None = None
                    _ocr_txt = os.environ.get(
                        "OMI_OCR_TEXT_DIR", str(self.workspace / "staging" / "ocr_text")
                    )
                    _report_d = Path(
                        os.environ.get("AIMS_REPORT_DIR", str(self.workspace / "report"))
                    )
                    for _search_dir in (
                        self.workspace / "inbox" / "Skipped",
                        self.workspace / "inbox" / "income",
                        self.workspace / "batch_inbox",
                        Path(_ocr_txt),
                        self.workspace / "outbox",
                        _report_d,
                        self.workspace / "inbox",
                        self.workspace / "master",
                    ):
                        _candidate = _search_dir / fn
                        if _candidate.is_file():
                            _fallback_path = _candidate
                            break
                    if _fallback_path is None:
                        miss.append(fn)
                        continue
                    # Build a fake_row from filesystem so AI helpers can run
                    _ext = _fallback_path.suffix.lower().lstrip(".")
                    _fake_row: dict = {
                        "file_name": fn,
                        "file_path": str(_fallback_path),
                        "title": "",
                        "summary": "",
                        "aims_process": "",
                    }
                    _sample = self._extract_text_sample(str(_fallback_path))
                    # Empty content → delete only if file is genuinely small (< 1 KB),
                    # not just because text extraction failed for this format
                    _file_size = _fallback_path.stat().st_size if _fallback_path.exists() else 0
                    if not (_sample or "").strip() and _file_size > 1024:
                        miss.append(fn)  # can't extract text but file has content — skip
                        continue
                    if not (_sample or "").strip():
                        if apply_changes:
                            try:
                                _fallback_path.unlink()
                                deleted.append(f"`{fn}` (пустое содержимое)" if lang == "ru" else f"`{fn}` (empty content)")
                            except Exception as _e:
                                fail.append(f"`{fn}`: delete error ({str(_e)[:120]})")
                        else:
                            deleted.append(f"`{fn}` (пустое содержимое — будет удалён)" if lang == "ru" else f"`{fn}` (empty content — would delete)")
                        continue
                    _ai_reason_fb = ""
                    if ai_mode:
                        _sug = None
                        if use_external_llm:
                            _sug = self._ai_rename_suggest_external(_fake_row, _sample)
                        if _sug is None:
                            _sug = self._ai_rename_suggest_qwen(_fake_row, _sample)
                        if _sug is not None:
                            _new_name_fb, _ai_reason_fb = _sug
                        else:
                            miss.append(fn)
                            continue
                    else:
                        # No AI — use sample first line as name hint
                        _first = (_sample or "").strip().split("\n")[0][:80].strip()
                        if _first:
                            import re as _re
                            _stem = _re.sub(r'[\\/:*?"<>|]', "_", _first).rstrip(".")[:80]
                            _new_name_fb = f"{_stem}.{_ext}" if _ext else _stem
                        else:
                            miss.append(fn)
                            continue
                    if not _new_name_fb or _new_name_fb == fn:
                        done.append(f"`{fn}` -> `{fn}`")
                        continue
                    if not apply_changes:
                        _suffix_fb = f" _(AI: {_ai_reason_fb})_" if _ai_reason_fb else ""
                        done.append(f"`{fn}` -> `{_new_name_fb}`{_suffix_fb}")
                        continue
                    _target_fb = _fallback_path.with_name(_new_name_fb)
                    if _target_fb.exists() and _target_fb.resolve() != _fallback_path.resolve():
                        # Duplicate name → delete source and report
                        if apply_changes:
                            try:
                                _fallback_path.unlink()
                                dupes.append(f"`{fn}` → уже есть `{_new_name_fb}`" if lang == "ru" else f"`{fn}` → already exists `{_new_name_fb}`")
                            except Exception as _e:
                                fail.append(f"`{fn}`: delete error ({str(_e)[:120]})")
                        else:
                            dupes.append(f"`{fn}` → уже есть `{_new_name_fb}` (будет удалён)" if lang == "ru" else f"`{fn}` → already exists `{_new_name_fb}` (would delete)")
                        continue
                    try:
                        if _target_fb.resolve() != _fallback_path.resolve():
                            _fallback_path.rename(_target_fb)
                        _suffix_fb = f" _(AI: {_ai_reason_fb})_" if _ai_reason_fb else ""
                        done.append(f"`{fn}` -> `{_new_name_fb}`{_suffix_fb}")
                    except Exception as _e:
                        fail.append(f"`{fn}`: rename error ({str(_e)[:120]})")
                    continue
                ai_reason = ""
                if ai_mode:
                    sample = self._extract_text_sample(str(row["file_path"] or ""))
                    suggested = None
                    if use_external_llm:
                        suggested = self._ai_rename_suggest_external(row, sample)
                    if suggested is None:
                        suggested = self._ai_rename_suggest_qwen(row, sample)
                    if suggested is not None:
                        new_name, ai_reason = suggested
                    else:
                        new_name = self._contextual_name_for_row(row)
                else:
                    new_name = self._contextual_name_for_row(row)
                old_name = str(row["file_name"] or "").strip()
                if not new_name or new_name == old_name:
                    done.append(f"`{old_name}` -> `{old_name}`")
                    continue
                if not apply_changes:
                    suffix = f" _(AI: {ai_reason})_" if ai_reason else ""
                    done.append(f"`{old_name}` -> `{new_name}`{suffix}")
                    continue
                fp = str(row["file_path"] or "").strip()
                if not fp:
                    fail.append(f"`{old_name}`: no file_path in registry")
                    continue
                p = Path(fp)
                if not p.is_file():
                    # File not at registered path — search all known pipeline locations
                    _ocr_txt_dir = Path(
                        os.environ.get(
                            "OMI_OCR_TEXT_DIR", str(self.workspace / "staging" / "ocr_text")
                        )
                    )
                    _search_locations = [
                        # 1. master (registered, anonymized, renamed)
                        self.workspace / "master",
                        # 2. Skipped (passed OCR but held)
                        self.workspace / "inbox" / "Skipped",
                        # 3. OCR sidecars + income (post-OCR / pre-quality-check)
                        _ocr_txt_dir,
                        self.workspace / "outbox",
                        self.workspace / "inbox" / "income",
                        # 4. batch_inbox (saved by Axi, awaiting Omi)
                        self.workspace / "batch_inbox",
                        self.workspace / "inbox",
                    ]
                    _found_p: Path | None = None
                    for _loc in _search_locations:
                        if not _loc.exists():
                            continue
                        for _candidate in _loc.rglob(old_name):
                            if _candidate.is_file():
                                _found_p = _candidate
                                break
                        if _found_p:
                            break
                    if _found_p:
                        # Found in alternate location — update DB path and continue rename
                        p = _found_p
                        fp = str(_found_p)
                        if apply_changes:
                            conn.execute(
                                "UPDATE documents SET file_path=? WHERE id=?",
                                (fp, row["id"]),
                            )
                    else:
                        # Genuinely gone — remove stale DB record
                        if apply_changes:
                            conn.execute("DELETE FROM documents WHERE id = ?", (row["id"],))
                            deleted.append(
                                f"`{old_name}` (файл не найден нигде, запись удалена из реестра)"
                                if lang == "ru" else
                                f"`{old_name}` (not found in any location, removed from registry)"
                            )
                        else:
                            deleted.append(
                                f"`{old_name}` (файл не найден нигде — будет удалён из реестра)"
                                if lang == "ru" else
                                f"`{old_name}` (not found in any location — would remove from registry)"
                            )
                        continue
                target = p.with_name(new_name)
                if target.exists() and target.resolve() != p.resolve():
                    # Duplicate name → delete source file and report
                    if apply_changes:
                        try:
                            p.unlink()
                            conn.execute("DELETE FROM documents WHERE id = ?", (row["id"],))
                            dupes.append(f"`{old_name}` → уже есть `{new_name}`" if lang == "ru" else f"`{old_name}` → already exists `{new_name}`")
                        except Exception as _e:
                            fail.append(f"`{old_name}`: delete error ({str(_e)[:120]})")
                    else:
                        dupes.append(f"`{old_name}` → уже есть `{new_name}` (будет удалён)" if lang == "ru" else f"`{old_name}` → already exists `{new_name}` (would delete)")
                    continue
                try:
                    if target.resolve() != p.resolve():
                        p.rename(target)
                    new_path = str(target)
                except Exception as e:
                    fail.append(f"`{old_name}`: rename error ({str(e)[:120]})")
                    continue
                conn.execute(
                    """
                    UPDATE documents
                       SET original_file_name = CASE
                                WHEN COALESCE(original_file_name, '') = '' THEN file_name
                                ELSE original_file_name
                           END,
                           canonical_file_name = ?,
                           file_name = ?,
                           file_path = ?,
                           date_modified = ?
                     WHERE id = ?
                    """,
                    (new_name, new_name, new_path, datetime.now().isoformat(), row["id"]),
                )
                suffix = f" _(AI: {ai_reason})_" if ai_reason else ""
                done.append(f"`{old_name}` -> `{new_name}`{suffix}")
            if apply_changes:
                conn.commit()

        if lang == "ru":
            title = "🧾 Dry-run переименования по контексту" if not apply_changes else "🧾 Переименование по контексту"
            lines = [f"{title}: *{len(done)}*"]
            lines.extend([f"- {x}" for x in done[:200]])
            if deleted:
                lines.append(f"\n🗑 Удалены (пустое содержимое): *{len(deleted)}*")
                lines.extend([f"- {x}" for x in deleted[:100]])
            if dupes:
                lines.append(f"\n⚠️ Удалены (дубликат уже есть): *{len(dupes)}*")
                lines.extend([f"- {x}" for x in dupes[:100]])
            if miss:
                lines.append("\nНе найдены в реестре:")
                lines.extend([f"- `{x}`" for x in miss[:100]])
            if fail and apply_changes:
                lines.append("\nОшибки (физический файл не переименован, реестр не изменён):")
                lines.extend([f"- {x}" for x in fail[:100]])
            if not apply_changes:
                lines.append("\nПрименить: `/rename_by_context apply file1.docx,file2.docx`")
            if ai_mode:
                lines.append("\nРежим: *AI rename* (Qwen локально" + (" + external fallback)." if use_external_llm else ")."))
            return "\n".join(lines)
        title = "🧾 Context rename dry-run" if not apply_changes else "🧾 Context rename done"
        lines = [f"{title}: *{len(done)}*"]
        lines.extend([f"- {x}" for x in done[:200]])
        if deleted:
            lines.append(f"\n🗑 Deleted (empty content): *{len(deleted)}*")
            lines.extend([f"- {x}" for x in deleted[:100]])
        if dupes:
            lines.append(f"\n⚠️ Deleted (duplicate already exists): *{len(dupes)}*")
            lines.extend([f"- {x}" for x in dupes[:100]])
        if miss:
            lines.append("\nNot found in registry:")
            lines.extend([f"- `{x}`" for x in miss[:100]])
        if fail and apply_changes:
            lines.append("\nErrors (physical file not renamed, registry unchanged):")
            lines.extend([f"- {x}" for x in fail[:100]])
        if not apply_changes:
            lines.append("\nApply: `/rename_by_context apply file1.docx,file2.docx`")
        if ai_mode:
            lines.append("\nMode: *AI rename* (local Qwen" + (" + external fallback)." if use_external_llm else ")."))
        return "\n".join(lines)

    @staticmethod
    def _normalize_registry_filename(name: str) -> str:
        """Collapse Unicode hyphens to ASCII and whitespace for loose matching."""
        s = (name or "").strip()
        for a, b in (("\u2013", "-"), ("\u2014", "-"), ("\u2010", "-"), ("\u00ad", "")):
            s = s.replace(a, b)
        s = re.sub(r"\s+", " ", s).strip()
        return s

    @staticmethod
    def _expand_registry_name_variants(name: str) -> list[str]:
        """
        Exact name, hyphen-normalized, and _anonymized / non-anonymized stem variants
        (batch/OCR often store different spellings than chat context).
        """
        raw = (name or "").strip()
        if not raw:
            return []
        seen: set[str] = set()
        out: list[str] = []

        def _push(s: str) -> None:
            s = s.strip()
            if s and s not in seen:
                seen.add(s)
                out.append(s)

        _push(raw)
        norm = StorageManager._normalize_registry_filename(raw)
        if norm != raw:
            _push(norm)
        p = Path(norm)
        stem, suf = p.stem, p.suffix
        suf_l = suf.lower()
        anon = "_anonymized"
        if stem.endswith(anon):
            _push(f"{stem[: -len(anon)]}{suf}")
        else:
            _push(f"{stem}{anon}{suf}")
        # OCR/registry often stores .txt while Telegram context still has .xlsx/.pdf upload name
        _sibling_ext = (".txt", ".xlsx", ".xls", ".pdf", ".docx", ".md", ".pptx", ".csv")
        for alt in _sibling_ext:
            if suf_l != alt:
                _push(f"{stem}{alt}")
        return out

    def _find_document_row_by_source_name(self, conn: sqlite3.Connection, src: str) -> sqlite3.Row | None:
        """Resolve documents row by file_name / source / canonical / original (exact + case-insensitive)."""
        variants = self._expand_registry_name_variants(src)
        if not variants:
            return None
        sql_select = """
                SELECT id, file_name, file_path,
                       COALESCE(source_filename, '') AS source_filename,
                       COALESCE(original_file_name, '') AS original_file_name,
                       COALESCE(canonical_file_name, '') AS canonical_file_name
                FROM documents
                """
        for c in variants:
            row = conn.execute(
                sql_select
                + """
                WHERE file_name = ?
                   OR COALESCE(source_filename, '') = ?
                   OR COALESCE(original_file_name, '') = ?
                   OR COALESCE(canonical_file_name, '') = ?
                ORDER BY COALESCE(date_added, date_modified, '') DESC
                LIMIT 1
                """,
                (c, c, c, c),
            ).fetchone()
            if row:
                return row
        for c in variants:
            row = conn.execute(
                sql_select
                + """
                WHERE lower(file_name) = lower(?)
                   OR lower(COALESCE(source_filename, '')) = lower(?)
                   OR lower(COALESCE(original_file_name, '')) = lower(?)
                   OR lower(COALESCE(canonical_file_name, '')) = lower(?)
                ORDER BY COALESCE(date_added, date_modified, '') DESC
                LIMIT 1
                """,
                (c, c, c, c),
            ).fetchone()
            if row:
                return row
        row = StorageManager._find_document_row_norm_sql(conn, src)
        if row:
            return row
        return StorageManager._find_document_row_like_unique(conn, src)

    @staticmethod
    def _sql_norm_lower_expr(column: str) -> str:
        """SQLite: normalize common Unicode dashes in a filename column, then lower()."""
        return (
            f"lower(replace(replace(replace(replace(replace({column}, char(8211), '-'), "
            f"char(8212), '-'), char(8210), '-'), char(8209), '-'), char(8722), '-'))"
        )

    @staticmethod
    def _find_document_row_norm_sql(conn: sqlite3.Connection, src: str) -> sqlite3.Row | None:
        """Match after Unicode-dash normalization on both DB side and search string."""
        variants = StorageManager._expand_registry_name_variants(src)
        if not variants:
            return None
        sql_select = """
                SELECT id, file_name, file_path,
                       COALESCE(source_filename, '') AS source_filename,
                       COALESCE(original_file_name, '') AS original_file_name,
                       COALESCE(canonical_file_name, '') AS canonical_file_name
                FROM documents
                """
        cols = (
            "file_name",
            "COALESCE(source_filename, '')",
            "COALESCE(original_file_name, '')",
            "COALESCE(canonical_file_name, '')",
        )
        for c in variants:
            n = StorageManager._normalize_registry_filename(c).lower()
            if not n:
                continue
            parts = []
            params: list[str] = []
            for col in cols:
                parts.append(f"{StorageManager._sql_norm_lower_expr(col)} = ?")
                params.append(n)
            row = conn.execute(
                sql_select + " WHERE " + " OR ".join(parts) + """
                ORDER BY COALESCE(date_added, date_modified, '') DESC
                LIMIT 1
                """,
                params,
            ).fetchone()
            if row:
                return row
        return None

    @staticmethod
    def _find_document_row_like_unique(conn: sqlite3.Connection, src: str) -> sqlite3.Row | None:
        """If exactly one row matches normalized stem (substring), use it."""
        n = StorageManager._normalize_registry_filename(src)
        p = Path(n)
        stem = p.stem
        anon = "_anonymized"
        if stem.lower().endswith(anon.lower()):
            stem = stem[: -len(anon)]
        stem = stem.strip()
        if len(stem) < 4:
            return None
        esc = stem.replace("\\", "\\\\").replace("%", "\\%").replace("_", "\\_")
        pattern = f"%{esc}%".lower()
        sql_select = """
                SELECT id, file_name, file_path,
                       COALESCE(source_filename, '') AS source_filename,
                       COALESCE(original_file_name, '') AS original_file_name,
                       COALESCE(canonical_file_name, '') AS canonical_file_name
                FROM documents
                """
        rows = conn.execute(
            sql_select
            + """
                WHERE lower(file_name) LIKE ? ESCAPE '\\'
                ORDER BY COALESCE(date_modified, date_added, '') DESC
                LIMIT 5
                """,
            (pattern,),
        ).fetchall()
        if len(rows) == 1:
            return rows[0]
        return None

    def _find_document_row_by_ocr_registry_id(self, conn: sqlite3.Connection, ocr_row_id: int) -> sqlite3.Row | None:
        """
        Telegram list #N sometimes matches ocr_documents.id, not documents.id.
        Resolve ocr row → aims documents by file_path / file_name.
        """
        if not self.ocr_db_path.is_file():
            return None
        try:
            oc = sqlite_connect_wal(self.ocr_db_path)
            oc.row_factory = sqlite3.Row
            r = oc.execute(
                "SELECT file_path, file_name FROM ocr_documents WHERE id = ?",
                (ocr_row_id,),
            ).fetchone()
            oc.close()
        except Exception:
            return None
        if not r:
            return None
        fp = str(r["file_path"] or "").strip()
        fn = str(r["file_name"] or "").strip()
        sql_sel = """
                SELECT id, file_name, file_path,
                       COALESCE(source_filename, '') AS source_filename,
                       COALESCE(original_file_name, '') AS original_file_name,
                       COALESCE(canonical_file_name, '') AS canonical_file_name
                FROM documents
                """
        for cand in ({fp, fp.replace("\\", "/")} if fp else set()):
            if not cand:
                continue
            row = conn.execute(sql_sel + " WHERE file_path = ? LIMIT 1", (cand,)).fetchone()
            if row:
                return row
        if fp:
            bn = Path(fp.replace("\\", "/")).name
            if bn:
                row = conn.execute(
                    sql_sel
                    + " WHERE lower(file_name) = lower(?) ORDER BY id DESC LIMIT 1",
                    (bn,),
                ).fetchone()
                if row:
                    return row
        if fn:
            return self._find_document_row_by_source_name(conn, fn)
        return None

    def rename_documents_by_request(
        self,
        source_name: str,
        requested_name: str,
        *,
        lang: str = "en",
        fallback_doc_id: int | None = None,
        source_document_id: int | None = None,
    ) -> str:
        """
        Rename one document to user-requested filename.
        Matches source by file_name/source_filename/original/canonical names,
        optional explicit documents.id (source_document_id), then chat fallback_doc_id.
        """
        lang = "ru" if (lang or "").lower().startswith("ru") else "en"
        src = (source_name or "").strip()
        req = (requested_name or "").strip()
        sid = int(source_document_id) if source_document_id is not None else 0
        if sid < 0:
            sid = 0
        # Fallback: id+name stayed in requested_name (e.g. client/parser never split "348_New name")
        if sid == 0 and req:
            _m = re.match(r"^#?(\d{1,8})(?:_|\s+)(.+)$", req, re.DOTALL)
            if _m and _m.group(2).strip():
                try:
                    _vid = int(_m.group(1))
                    _digits = _m.group(1)
                    # avoid "2024_report.xlsx" → id=2024; allow 348, 12345, etc.
                    _looks_like_year = len(_digits) == 4 and 1990 <= _vid <= 2035
                    if not _looks_like_year:
                        sid = _vid
                        req = _m.group(2).strip()
                except ValueError:
                    pass
        if not req:
            return "⚠️ Укажи новое имя файла." if lang == "ru" else "⚠️ Provide target filename."
        if re.fullmatch(r"#?\d+", req):
            return (
                "⚠️ Укажите *новое имя* после id: `/rename_file_348_новое_имя.xlsx` "
                "(или реплай + `/rename_file_новое_имя.xlsx`)."
                if lang == "ru"
                else "⚠️ Put the *new name* after the id: `/rename_file_348_new_name.xlsx` "
                "(or reply to the file + `/rename_file_newname.xlsx`)."
            )
        if not src and not sid:
            return (
                "⚠️ Неизвестен файл: ответьте реплаем на документ или используйте "
                "`/rename_file_<id>_<новое_имя>` (например `/rename_file_348_Plan.xlsx`)."
                if lang == "ru"
                else "⚠️ Unknown file: reply to the document or use "
                "`/rename_file_<id>_<new_name>` (e.g. `/rename_file_348_Plan.xlsx`)."
            )

        with self._conn() as conn:
            row = None
            if sid > 0:
                row = conn.execute(
                    """
                    SELECT id, file_name, file_path,
                           COALESCE(source_filename, '') AS source_filename,
                           COALESCE(original_file_name, '') AS original_file_name,
                           COALESCE(canonical_file_name, '') AS canonical_file_name
                    FROM documents WHERE id = ?
                    """,
                    (sid,),
                ).fetchone()
            if not row and sid > 0:
                row = self._find_document_row_by_ocr_registry_id(conn, sid)
            if not row and src:
                row = self._find_document_row_by_source_name(conn, src)
            if not row and fallback_doc_id is not None and fallback_doc_id > 0:
                row = conn.execute(
                    """
                    SELECT id, file_name, file_path,
                           COALESCE(source_filename, '') AS source_filename,
                           COALESCE(original_file_name, '') AS original_file_name,
                           COALESCE(canonical_file_name, '') AS canonical_file_name
                    FROM documents WHERE id = ?
                    """,
                    (fallback_doc_id,),
                ).fetchone()
            if not row:
                if sid > 0:
                    return (
                        f"⚠️ Номер `#{sid}` не найден: нет строки `documents.id`, "
                        f"и нет совпадения по `ocr_documents` → `documents` "
                        f"(разные id в OCR и в AIMS). Проверьте `#` в том же списке или синк OCR→AIMS."
                        if lang == "ru"
                        else f"⚠️ id `{sid}` not found: no `documents.id`, and no "
                        f"`ocr_documents`→`documents` match (OCR vs AIMS ids may differ). "
                        f"Confirm `#` from the latest list or run OCR→AIMS sync."
                    )
                return (
                    f"⚠️ Файл `{src}` не найден в реестре."
                    if lang == "ru"
                    else f"⚠️ File `{src}` not found in registry."
                )

            old_name = str(row["file_name"] or "").strip() or src
            file_path = str(row["file_path"] or "").strip()
            p = Path(file_path) if file_path else Path()
            ext_old = Path(old_name).suffix
            req_path = Path(req)
            ext_new = req_path.suffix or ext_old
            stem = _sanitize_filename_stem(req_path.stem or req_path.name)
            new_name = f"{stem}{ext_new}" if ext_new else stem
            if new_name == old_name:
                return (
                    f"ℹ️ Имя уже актуально: `{old_name}`"
                    if lang == "ru"
                    else f"ℹ️ Name already set: `{old_name}`"
                )

            if not p.is_file():
                return (
                    f"⚠️ Файл `{old_name}` отсутствует на диске, переименование невозможно."
                    if lang == "ru"
                    else f"⚠️ File `{old_name}` is missing on disk; cannot rename."
                )

            target = p.with_name(new_name)
            if target.exists() and target.resolve() != p.resolve():
                return (
                    f"⚠️ Целевое имя уже занято: `{new_name}`."
                    if lang == "ru"
                    else f"⚠️ Target filename already exists: `{new_name}`."
                )
            try:
                if target.resolve() != p.resolve():
                    p.rename(target)
            except Exception as e:
                return (
                    f"❌ Ошибка переименования: {type(e).__name__}"
                    if lang == "ru"
                    else f"❌ Rename failed: {type(e).__name__}"
                )

            conn.execute(
                """
                UPDATE documents
                   SET original_file_name = CASE
                            WHEN COALESCE(original_file_name, '') = '' THEN file_name
                            ELSE original_file_name
                       END,
                       canonical_file_name = ?,
                       file_name = ?,
                       file_path = ?,
                       date_modified = ?
                 WHERE id = ?
                """,
                (new_name, new_name, str(target), datetime.now().isoformat(), row["id"]),
            )
            conn.commit()

        return (
            f"✅ Переименовано по запросу: `{old_name}` → `{new_name}`"
            if lang == "ru"
            else f"✅ Renamed by request: `{old_name}` → `{new_name}`"
        )

    # ── Batch auto-fix for noisy (numeric-code-laden) filenames ───────────────

    _FIX_NOISY_HIGH = re.compile(r"\d{6,}")
    _FIX_NOISY_STD = re.compile(
        r"(?:ISO|IEC|API|ANSI|ASME|ASTM|NFPA|IEEE|EN|DIN|BS|GOST)\s*\d{1,5}",
        re.IGNORECASE,
    )

    def _strip_noisy_codes(self, name: str) -> str:
        """Remove 6+ digit numeric codes from filename stem, preserve extension."""
        ext = Path(name).suffix
        stem = Path(name).stem
        protected = self._FIX_NOISY_STD.sub("\x00STD\x00", stem)
        # Remove 6+ digit runs (with optional surrounding separators)
        cleaned = re.sub(r"[\s_\-]*\d{6,}[\s_\-]*", " ", protected)
        cleaned = cleaned.replace("\x00STD\x00", "STD")
        cleaned = re.sub(r"[\s_]+", " ", cleaned).strip().strip("_-.")
        if not cleaned:
            return ""
        return f"{cleaned}{ext}"

    def fix_noisy_names_batch(
        self,
        limit: int = 500,
        apply_changes: bool = False,
        lang: str = "ru",
    ) -> str:
        """
        Scan registry for files whose names contain long numeric codes (6+ digits),
        auto-strip the codes, and optionally rename them in DB + on disk.
        """
        lang = "ru" if (lang or "").lower().startswith("ru") else "en"
        with self._conn() as conn:
            rows = conn.execute(
                "SELECT id, file_name FROM documents ORDER BY id DESC LIMIT ?",
                (limit,),
            ).fetchall()

        candidates: list[tuple[int, str, str]] = []
        for doc_id, file_name in rows:
            name = str(file_name or "").strip()
            if not name:
                continue
            stem = Path(name).stem
            protected_stem = self._FIX_NOISY_STD.sub("STD", stem)
            if not self._FIX_NOISY_HIGH.search(protected_stem):
                continue
            new_name = self._strip_noisy_codes(name)
            if not new_name or new_name == name:
                continue
            candidates.append((doc_id, name, new_name))

        if not candidates:
            return (
                "✅ Шумовых числовых кодов в именах файлов не найдено."
                if lang == "ru"
                else "✅ No noisy numeric codes found in filenames."
            )

        if not apply_changes:
            lines = [
                f"🔍 {'Найдено' if lang == 'ru' else 'Found'} {len(candidates)} "
                f"{'файлов с шумовыми именами' if lang == 'ru' else 'files with noisy names'} "
                f"({'dry-run'}):"
            ]
            for doc_id, old, new in candidates[:40]:
                lines.append(f"  #{doc_id}  `{old}`\n       → `{new}`")
            if len(candidates) > 40:
                lines.append(
                    f"  ... {'и ещё' if lang == 'ru' else 'and'} {len(candidates) - 40} "
                    f"{'файлов' if lang == 'ru' else 'more'}"
                )
            lines.append(
                f"\n{'Для применения' if lang == 'ru' else 'To apply'}: `/fixnoisynames apply`"
            )
            return "\n".join(lines)

        done: list[str] = []
        fail: list[str] = []
        for doc_id, old_name, new_name in candidates:
            result = self.rename_documents_by_request(
                source_name="",
                requested_name=new_name,
                lang=lang,
                source_document_id=doc_id,
            )
            if result.startswith("✅"):
                done.append(f"  #{doc_id}: `{old_name}` → `{new_name}`")
            else:
                fail.append(f"  #{doc_id}: {result}")

        lines = []
        if done:
            hdr = f"✅ Переименовано {len(done)}:" if lang == "ru" else f"✅ Renamed {len(done)}:"
            lines.append(hdr)
            lines.extend(done[:30])
            if len(done) > 30:
                lines.append(f"  ... {'и ещё' if lang == 'ru' else 'and'} {len(done) - 30}")
        if fail:
            hdr = f"❌ Ошибки {len(fail)}:" if lang == "ru" else f"❌ Errors {len(fail)}:"
            lines.append(hdr)
            lines.extend(fail[:10])
        return "\n".join(lines) or ("Ничего не изменено." if lang == "ru" else "Nothing changed.")

    def list_documents_today(self, limit: int = 200, lang: str = "en") -> str:
        """
        Список документов, обработанных сегодня (00:00–23:59 локального времени).
        Опирается на date_added, с fallback на date_modified.
        """
        lang = "ru" if (lang or "").lower().startswith("ru") else "en"
        today = datetime.now().strftime("%Y-%m-%d")
        with self._conn() as conn:
            rows = conn.execute(
                """
                SELECT
                    id,
                    file_name,
                    file_path,
                    COALESCE(file_created_at, '') AS fca,
                    COALESCE(file_mtime_at, '') AS fma,
                    COALESCE(date_added, '') AS da,
                    COALESCE(date_modified, '') AS dm,
                    COALESCE(aims_process, '') AS proc,
                    COALESCE(source, '') AS src
                FROM documents
                WHERE (date_added LIKE ? OR date_modified LIKE ?)
                ORDER BY COALESCE(date_added, date_modified, '') DESC
                LIMIT ?
                """,
                (f"{today}%", f"{today}%", int(limit)),
            ).fetchall()

        if not rows:
            if lang == "ru":
                return "📄 За сегодня (00:00–23:59) обработанных документов нет."
            return "📄 No documents processed today (00:00–23:59)."

        if lang == "ru":
            lines = [f"📄 Документы за сегодня ({today}, 00:00–23:59): *{len(rows)}*"]
            na = "(без имени)"
            reg_l, fs_l = "внесено в реестр", "файл (mtime)"
        else:
            lines = [f"📄 Documents today ({today}, 00:00–23:59): *{len(rows)}*"]
            na = "(no name)"
            reg_l, fs_l = "registered", "file mtime"
        for r in rows:
            fn = str(r["file_name"] or "").strip() or na
            reg = str(r["da"] or "").replace("T", " ")[:19] or "—"
            fs = str(r["fma"] or "").replace("T", " ")[:19] or _file_mtime_short(str(r["file_path"] or "")) or "—"
            proc = str(r["proc"] or "").strip()
            src = str(r["src"] or "").strip()
            meta = []
            if proc:
                meta.append(proc)
            if src:
                meta.append(src)
            meta_s = f" — {', '.join(meta)}" if meta else ""
            lines.append(
                f"- #{r['id']} `{fn}`{meta_s}\n"
                f"  · {reg_l}: {reg} · {fs_l}: {fs}"
            )
        return "\n".join(lines)

    def list_documents_last_hours(self, hours: int = 24, limit: int = 300, lang: str = "en") -> str:
        """
        Список документов с активностью в реестре за последние N часов.
        Учитывается max(date_added, date_modified) по парсящимся меткам (см. _latest_registry_activity_ts).
        """
        lang = "ru" if (lang or "").lower().startswith("ru") else "en"
        try:
            hrs = max(1, min(24 * 30, int(hours)))
        except Exception:
            hrs = 24
        now = datetime.now()
        prelimit = max(400, int(limit) * 4)
        with self._conn() as conn:
            rows = conn.execute(
                """
                SELECT
                    id,
                    file_name,
                    file_path,
                    COALESCE(file_created_at, '') AS fca,
                    COALESCE(file_mtime_at, '') AS fma,
                    COALESCE(date_added, date_modified, '') AS dt,
                    COALESCE(date_added, '') AS da,
                    COALESCE(date_modified, '') AS dm,
                    COALESCE(aims_process, '') AS proc,
                    COALESCE(source, '') AS src
                FROM documents
                ORDER BY COALESCE(NULLIF(trim(date_modified), ''), NULLIF(trim(date_added), ''), '') DESC
                LIMIT ?
                """,
                (prelimit,),
            ).fetchall()

        cutoff = now.timestamp() - (hrs * 3600)
        out = []
        for r in rows:
            da = str(r["da"] or "").strip()
            dm = str(r["dm"] or "").strip()
            ts = _latest_registry_activity_ts(da, dm)
            if ts is None or ts < cutoff:
                continue
            out.append(r)
            if len(out) >= int(limit):
                break

        if not out:
            if lang == "ru":
                return f"📄 За последние {hrs} ч в реестре новых записей нет."
            return f"📄 No registry entries in the last {hrs}h."

        if lang == "ru":
            lines = [f"📄 За последние {hrs} ч (по max дате внесения/обновления в БД): *{len(out)}*"]
            na = "(без имени)"
            reg_l, fs_l, upd_l = "внесено в реестр", "файл (mtime)", "обновлено в БД"
        else:
            lines = [f"📄 Last {hrs}h (by *latest of added/updated* in DB): *{len(out)}*"]
            na = "(no name)"
            reg_l, fs_l, upd_l = "registered in DB", "file mtime", "DB updated"
        for r in out:
            fn = str(r["file_name"] or "").strip() or na
            reg = str(r["da"] or "").replace("T", " ")[:19] or str(r["dt"] or "").replace("T", " ")[:19]
            dm = str(r["dm"] or "").replace("T", " ")[:19]
            fs = (
                str(r["fma"] or "").replace("T", " ")[:19]
                or _file_mtime_short(str(r["file_path"] or ""))
                or ("—" if lang == "ru" else "n/a")
            )
            proc = str(r["proc"] or "").strip()
            src = str(r["src"] or "").strip()
            meta = []
            if proc:
                meta.append(proc)
            if src:
                meta.append(src)
            meta_s = f" — {', '.join(meta)}" if meta else ""
            upd_part = ""
            if dm and dm != reg[:19]:
                upd_part = f" · {upd_l}: {dm}"
            lines.append(
                f"- #{r['id']} `{fn}`{meta_s}\n"
                f"  · {reg_l}: {reg} · {fs_l}: {fs}{upd_part}"
            )
        return "\n".join(lines)

    def list_documents_by_process_context(
        self, hours: int = 24, limit: int = 400, lang: str = "en"
    ) -> str:
        """
        Краткая сводка документов по процессу за окно по часам.
        Окно считается по max(date_added, date_modified) — как в list_documents_last_hours.
        """
        lang = "ru" if (lang or "").lower().startswith("ru") else "en"
        try:
            hrs = max(1, min(24 * 30, int(hours)))
        except Exception:
            hrs = 24
        now = datetime.now()
        prelimit = max(400, int(limit) * 4)
        with self._conn() as conn:
            rows = conn.execute(
                """
                SELECT
                    file_name,
                    COALESCE(date_added, '') AS da,
                    COALESCE(date_modified, '') AS dm,
                    COALESCE(aims_process, 'UNASSIGNED') AS proc
                FROM documents
                ORDER BY COALESCE(NULLIF(trim(date_modified), ''), NULLIF(trim(date_added), ''), '') DESC
                LIMIT ?
                """,
                (prelimit,),
            ).fetchall()

        cutoff = now.timestamp() - (hrs * 3600)
        grouped: dict[str, list[str]] = {}
        for r in rows:
            ts = _latest_registry_activity_ts(str(r["da"] or ""), str(r["dm"] or ""))
            if ts is None or ts < cutoff:
                continue
            proc = str(r["proc"] or "").strip() or "UNASSIGNED"
            grouped.setdefault(proc, []).append(str(r["file_name"] or "").strip() or "(no name)")

        if not grouped:
            if lang == "ru":
                return f"📄 За последние {hrs} ч по контекстам записей нет."
            return f"📄 No registry entries by context in the last {hrs}h."

        total = sum(len(v) for v in grouped.values())
        if lang == "ru":
            lines = [f"📄 По контексту (за последние {hrs} ч): *{total}*"]
            more = "… ещё"
        else:
            lines = [f"📄 By process context (last {hrs}h): *{total}*"]
            more = "... and"
        for proc in sorted(grouped.keys()):
            files = grouped[proc]
            unit = "файл(ов)" if lang == "ru" else "file(s)"
            lines.append(f"\n*{proc}* — {len(files)} {unit}")
            for fn in files[:50]:
                lines.append(f"- `{fn}`")
            if len(files) > 50:
                if lang == "ru":
                    lines.append(f"- {more} {len(files) - 50}")
                else:
                    lines.append(f"- {more} {len(files) - 50} more")
        return "\n".join(lines)

    def list_registry_snapshot(self, limit: int = 60, lang: str = "en") -> str:
        """
        Последние записи реестра (без окна по часам) — для запросов «покажи реестр» без периода.
        Показывает дату внесения в БД и mtime файла на диске.
        """
        lang = "ru" if (lang or "").lower().startswith("ru") else "en"
        lim = max(1, min(500, int(limit)))
        with self._conn() as conn:
            rows = conn.execute(
                """
                SELECT
                    id,
                    file_name,
                    file_path,
                    COALESCE(file_created_at, '') AS fca,
                    COALESCE(file_mtime_at, '') AS fma,
                    COALESCE(date_added, '') AS da,
                    COALESCE(date_modified, '') AS dm,
                    COALESCE(aims_process, '') AS proc,
                    COALESCE(source, '') AS src
                FROM documents
                ORDER BY COALESCE(date_added, date_modified, '') DESC
                LIMIT ?
                """,
                (lim,),
            ).fetchall()

        if not rows:
            if lang == "ru":
                return "📄 Реестр документов пуст."
            return "📄 The document registry is empty."

        if lang == "ru":
            lines = [f"📄 Последние записи реестра (до *{len(rows)}*):\n"]
            na = "(без имени)"
            reg_l, fs_l, upd_l = "внесено в реестр", "файл (mtime)", "обновлено в БД"
        else:
            lines = [f"📄 Latest registry entries (up to *{len(rows)}*):\n"]
            na = "(no name)"
            reg_l, fs_l, upd_l = "registered in DB", "file mtime", "DB updated"
        for r in rows:
            fn = str(r["file_name"] or "").strip() or na
            reg = str(r["da"] or "").replace("T", " ")[:19] or "—"
            dm = str(r["dm"] or "").replace("T", " ")[:19]
            fs = (
                str(r["fma"] or "").replace("T", " ")[:19]
                or _file_mtime_short(str(r["file_path"] or ""))
                or ("—" if lang == "ru" else "n/a")
            )
            proc = str(r["proc"] or "").strip()
            src = str(r["src"] or "").strip()
            meta = []
            if proc:
                meta.append(proc)
            if src:
                meta.append(src)
            meta_s = f" — {', '.join(meta)}" if meta else ""
            upd_part = ""
            if dm and dm != reg and reg != "—":
                upd_part = f" · {upd_l}: {dm}"
            elif dm and reg == "—":
                upd_part = f" · {upd_l}: {dm}"
            lines.append(
                f"- #{r['id']} `{fn}`{meta_s}\n"
                f"  · {reg_l}: {reg} · {fs_l}: {fs}{upd_part}"
            )
        return "\n".join(lines)

    def list_documents_between(
        self,
        start: datetime,
        end: datetime,
        *,
        limit: int = 400,
        lang: str = "en",
    ) -> str:
        """
        Записи, у которых дата внесения **или** последнего обновления (по префиксу YYYY-MM-DD)
        попадает в календарный диапазон.
        """
        lang = "ru" if (lang or "").lower().startswith("ru") else "en"
        if start > end:
            start, end = end, start
        d1 = start.strftime("%Y-%m-%d")
        d2 = end.strftime("%Y-%m-%d")
        lim = max(1, min(2000, int(limit)))
        with self._conn() as conn:
            rows = conn.execute(
                """
                SELECT
                    id,
                    file_name,
                    file_path,
                    COALESCE(file_created_at, '') AS fca,
                    COALESCE(file_mtime_at, '') AS fma,
                    COALESCE(date_added, '') AS da,
                    COALESCE(date_modified, '') AS dm,
                    COALESCE(aims_process, '') AS proc,
                    COALESCE(source, '') AS src
                FROM documents
                WHERE (
                    (length(trim(COALESCE(date_added, ''))) >= 10
                     AND substr(trim(date_added), 1, 10) BETWEEN ? AND ?)
                    OR
                    (length(trim(COALESCE(date_modified, ''))) >= 10
                     AND substr(trim(date_modified), 1, 10) BETWEEN ? AND ?)
                )
                ORDER BY COALESCE(NULLIF(trim(date_modified), ''), NULLIF(trim(date_added), ''), '') DESC
                LIMIT ?
                """,
                (d1, d2, d1, d2, lim),
            ).fetchall()

        if not rows:
            if lang == "ru":
                return f"📄 За период *{d1}* … *{d2}* записей в реестре нет."
            return f"📄 No registry entries between *{d1}* and *{d2}*."

        if lang == "ru":
            lines = [f"📄 Реестр за *{d1}* … *{d2}*: *{len(rows)}*"]
            na = "(без имени)"
            reg_l, fs_l, upd_l = "внесено в реестр", "файл (mtime)", "обновлено в БД"
        else:
            lines = [f"📄 Registry *{d1}* … *{d2}*: *{len(rows)}*"]
            na = "(no name)"
            reg_l, fs_l, upd_l = "registered in DB", "file mtime", "DB updated"
        for r in rows:
            fn = str(r["file_name"] or "").strip() or na
            reg = str(r["da"] or "").replace("T", " ")[:19] or "—"
            dm = str(r["dm"] or "").replace("T", " ")[:19]
            fs = (
                str(r["fma"] or "").replace("T", " ")[:19]
                or _file_mtime_short(str(r["file_path"] or ""))
                or ("—" if lang == "ru" else "n/a")
            )
            proc = str(r["proc"] or "").strip()
            src = str(r["src"] or "").strip()
            meta = []
            if proc:
                meta.append(proc)
            if src:
                meta.append(src)
            meta_s = f" — {', '.join(meta)}" if meta else ""
            upd_part = ""
            if dm and dm != reg and reg != "—":
                upd_part = f" · {upd_l}: {dm}"
            elif dm and reg == "—":
                upd_part = f" · {upd_l}: {dm}"
            lines.append(
                f"- #{r['id']} `{fn}`{meta_s}\n"
                f"  · {reg_l}: {reg} · {fs_l}: {fs}{upd_part}"
            )
        return "\n".join(lines)

    def candidate_paths_for_rag(
        self,
        user_message: str,
        *,
        limit: int = 12,
    ) -> list[str]:
        """
        Локальные пути к файлам для qwen-agent RAG (retrieval).
        Предфильтрация по ключевым словам из сообщения + недавние документы,
        чтобы не парсить весь реестр на каждый запрос.

        Расширения должны совпадать с `qwen_agent.tools.simple_doc_parser.PARSER_SUPPORTED_FILE_TYPES`.
        """
        # keep in sync with qwen_agent (pdf, docx, pptx, txt, html, csv, tsv, xlsx, xls)
        allowed_ext = frozenset(
            ("pdf", "docx", "pptx", "txt", "html", "csv", "tsv", "xlsx", "xls")
        )
        ws = self.workspace.resolve()
        raw = (user_message or "").strip()
        tokens = [t for t in re.split(r"\W+", raw.lower()) if len(t) >= 3][:10]

        with self._conn() as conn:
            if len(tokens) >= 2:
                parts: list[str] = []
                params: list = []
                for _t in tokens:
                    parts.append(
                        "(LOWER(COALESCE(title,'')) LIKE ? OR LOWER(COALESCE(summary,'')) LIKE ? "
                        "OR LOWER(COALESCE(keywords,'')) LIKE ? OR LOWER(file_name) LIKE ?)"
                    )
                    p = f"%{_t}%"
                    params.extend([p, p, p, p])
                sql = (
                    "SELECT file_path FROM documents WHERE "
                    + " OR ".join(parts)
                    + " ORDER BY date_added DESC LIMIT ?"
                )
                rows = conn.execute(sql, params + [limit * 4]).fetchall()
            else:
                rows = conn.execute(
                    "SELECT file_path FROM documents ORDER BY date_added DESC LIMIT ?",
                    (limit * 4,),
                ).fetchall()

        out: list[str] = []
        for r in rows:
            fp = Path(r["file_path"])
            if not fp.is_file():
                continue
            try:
                fp.resolve().relative_to(ws)
            except ValueError:
                continue
            ext = fp.suffix.lower().lstrip(".")
            if ext not in allowed_ext:
                continue
            s = str(fp.resolve())
            if s not in out:
                out.append(s)
            if len(out) >= limit:
                break
        return out

    # ── 3. Переместить файлы ──────────────────────────────────

    def move_documents(self, query: str, target_process: str) -> str:
        # Проверить что процесс существует
        with self._conn() as conn:
            proc = conn.execute(
                "SELECT code, name, path FROM omi_processes WHERE code=?", (target_process,)
            ).fetchone()

        if not proc:
            return (f"❌ Процесс `{target_process}` не найден.\n"
                    f"Доступные: {', '.join(p['code'] for p in self.list_processes())}")

        # Найти файлы
        q = f"%{query}%"
        with self._conn() as conn:
            rows = conn.execute(
                "SELECT id, file_path, file_name, aims_process FROM documents "
                "WHERE file_name LIKE ? OR title LIKE ? OR keywords LIKE ?",
                (q, q, q)
            ).fetchall()

        if not rows:
            return f"❌ Файлы по запросу *«{query}»* не найдены."

        target_dir = Path(proc["path"])
        target_dir.mkdir(parents=True, exist_ok=True)

        moved = []
        errors = []
        for row in rows:
            src = Path(row["file_path"])
            dst = target_dir / row["file_name"]
            try:
                if src.exists():
                    shutil.move(str(src), str(dst))
                # Обновить БД
                with self._conn() as conn:
                    conn.execute(
                        "UPDATE documents SET file_path=?, aims_process=?, date_modified=? WHERE id=?",
                        (str(dst), target_process, datetime.now().isoformat(), row["id"])
                    )
                moved.append(row["file_name"])
            except Exception as e:
                errors.append(f"{row['file_name']}: {e}")

        lines = [f"📁 *Перемещено в `{target_process}` ({proc['name']}):*\n"]
        for f in moved:
            lines.append(f"  ✓ {f}")
        if errors:
            lines.append("\n⚠️ *Ошибки:*")
            for e in errors:
                lines.append(f"  ✗ {e}")
        return "\n".join(lines)

    # ── 4. Архивирование ─────────────────────────────────────

    def archive_documents(self, process: str = None, before_year: int = None) -> str:
        archive_dir = self.workspace / "archive"
        archive_dir.mkdir(parents=True, exist_ok=True)

        sql = "SELECT id, file_path, file_name, aims_process, date_added FROM documents WHERE 1=1"
        params = []

        if process:
            sql += " AND aims_process = ?"
            params.append(process)

        if before_year:
            sql += " AND date_added < ?"
            params.append(f"{before_year}-01-01")

        with self._conn() as conn:
            rows = conn.execute(sql, params).fetchall()

        if not rows:
            return "📦 Нет документов для архивирования по заданным критериям."

        archived = []
        for row in rows:
            src = Path(row["file_path"])
            year_dir = archive_dir / (row["aims_process"] or "unknown") / \
                       (row["date_added"] or "")[:4]
            year_dir.mkdir(parents=True, exist_ok=True)
            dst = year_dir / row["file_name"]
            try:
                if src.exists():
                    shutil.move(str(src), str(dst))
                with self._conn() as conn:
                    conn.execute(
                        "UPDATE documents SET file_path=?, notes=?, date_modified=? WHERE id=?",
                        (str(dst), "ARCHIVED", datetime.now().isoformat(), row["id"])
                    )
                archived.append(row["file_name"])
            except Exception as e:
                archived.append(f"⚠️ {row['file_name']}: {e}")

        lines = [f"🗄 *Архивировано {len(archived)} файлов:*\n"]
        for f in archived[:15]:
            lines.append(f"  ✓ {f}")
        if len(archived) > 15:
            lines.append(f"  ... и ещё {len(archived)-15}")
        lines.append(f"\n📂 Архив: `{archive_dir}`")
        return "\n".join(lines)

    # ── 5. Новый процесс ─────────────────────────────────────

    def create_process(self, code: str, name: str) -> str:
        folder_name = f"{code}_{name.replace(' ', '_')}"
        proc_dir = self.workspace / "master" / folder_name
        try:
            proc_dir.mkdir(parents=True, exist_ok=True)
            with self._conn() as conn:
                conn.execute(
                    "INSERT OR REPLACE INTO omi_processes(code,name,path,created) VALUES(?,?,?,?)",
                    (code, name, str(proc_dir), datetime.now().isoformat())
                )
            return (f"✅ *Новый раздел создан:*\n\n"
                    f"  Код: `{code}`\n"
                    f"  Название: *{name}*\n"
                    f"  Папка: `{proc_dir}`\n\n"
                    f"Документы можно добавлять командой:\n"
                    f"`/move <файл> {code}`")
        except Exception as e:
            return f"❌ Ошибка создания раздела: {e}"

    # ── 6. Миграция БД ───────────────────────────────────────

    def migrate_db(self, new_path_str: str) -> str:
        new_path = Path(new_path_str).expanduser()
        if new_path == self.db_path:
            return "ℹ️ БД уже находится по этому пути."

        try:
            new_path.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(str(self.db_path), str(new_path))

            # Обновить конфиг
            with self._conn() as conn:
                conn.execute(
                    "INSERT OR REPLACE INTO omi_config(key,value,updated) VALUES(?,?,?)",
                    ("db_path", str(new_path), datetime.now().isoformat())
                )

            return (f"✅ *БД скопирована:*\n\n"
                    f"  Откуда: `{self.db_path}`\n"
                    f"  Куда:   `{new_path}`\n\n"
                    f"⚠️ Обнови `OMI_DB_PATH` в `.env` и перезапусти Omi:\n"
                    f"`OMI_DB_PATH={new_path}`")
        except Exception as e:
            return f"❌ Ошибка миграции БД: {e}"

    def get_dgx_migration_plan(self) -> str:
        return (
            "🚀 *Миграция на Nvidia DGX Spark*\n\n"
            "*Шаг 1 — Скопировать БД на DGX:*\n"
            "```\n"
            f"rsync -av --progress \\\n"
            f"  {self.db_path} \\\n"
            f"  dgx:/home/axiomsphere/aims-data/\n"
            "```\n\n"
            "*Шаг 2 — Скопировать файлы:*\n"
            "```\n"
            f"rsync -av --progress \\\n"
            f"  {self.workspace}/ \\\n"
            f"  dgx:/aims_workspace/\n"
            "```\n\n"
            "*Шаг 3 — На DGX обновить .env:*\n"
            "```\n"
            "OMI_DB_PATH=/home/axiomsphere/aims-data/omi_registry.db\n"
            "AIMS_WORKSPACE=/aims_workspace\n"
            "```\n\n"
            "*Шаг 4 — Запустить Omi на DGX:*\n"
            "```\n"
            "python omi_bot.py\n"
            "```\n\n"
            "✅ Конвертация не нужна. SQLite — один файл."
        )

    # ── Утилиты ───────────────────────────────────────────────

    def _detect_platform(self) -> str:
        system = platform.system()
        if system == "Linux":
            # Проверить WSL
            try:
                with open("/proc/version") as f:
                    if "microsoft" in f.read().lower():
                        return "WSL (Ubuntu on Windows)"
            except Exception:
                pass
            return "Linux"
        elif system == "Windows":
            return "Windows"
        elif system == "Darwin":
            return "macOS"
        return system

    # ── Entity memory (chat context) ──────────────────────────

    def get_chat_context(self, chat_id: int) -> dict:
        """Return last known context for this chat (last file, process, doc_type, search)."""
        try:
            with self._conn() as conn:
                row = conn.execute(
                    "SELECT * FROM omi_chat_context WHERE chat_id = ?", (chat_id,)
                ).fetchone()
                if row:
                    return dict(row)
        except Exception:
            pass
        return {}

    def set_chat_context(
        self,
        chat_id: int,
        *,
        last_file_name: str | None = None,
        last_file_path: str | None = None,
        last_process: str | None = None,
        last_doc_type: str | None = None,
        last_search: str | None = None,
        last_doc_id: int | None = None,
    ) -> None:
        """Upsert entity context for a chat — only update provided fields."""
        try:
            now = datetime.now().isoformat()
            with self._conn() as conn:
                existing = conn.execute(
                    "SELECT * FROM omi_chat_context WHERE chat_id = ?", (chat_id,)
                ).fetchone()
                if existing:
                    updates, params = [], []
                    for col, val in [
                        ("last_file_name", last_file_name),
                        ("last_file_path", last_file_path),
                        ("last_process", last_process),
                        ("last_doc_type", last_doc_type),
                        ("last_search", last_search),
                        ("last_doc_id", last_doc_id),
                    ]:
                        if val is not None:
                            updates.append(f"{col} = ?")
                            params.append(val)
                    if updates:
                        params += [now, chat_id]
                        conn.execute(
                            f"UPDATE omi_chat_context SET {', '.join(updates)}, updated_at = ? WHERE chat_id = ?",
                            params,
                        )
                else:
                    conn.execute(
                        """INSERT INTO omi_chat_context
                           (chat_id, last_file_name, last_file_path, last_process,
                            last_doc_type, last_search, last_doc_id, updated_at)
                           VALUES (?,?,?,?,?,?,?,?)""",
                        (chat_id, last_file_name, last_file_path, last_process,
                         last_doc_type, last_search, last_doc_id, now),
                    )
                conn.commit()
        except Exception:
            pass

    # ── Soft-scored search for file delivery ─────────────────

    def search_document_files_scored(
        self,
        query: str,
        *,
        process: str | None = None,
        is_master: bool | None = None,
        year: str | None = None,
        doc_type: str | None = None,
        limit: int = 10,
    ) -> list[dict]:
        """Back-compat alias: use hybrid rerank."""
        return self.hybrid_search_documents(
            query,
            process=process,
            is_master=is_master,
            year=year,
            doc_type=doc_type,
            limit=limit,
        )
